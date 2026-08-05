"""Smoke-check catalog formulas, readings, worked lines, and answers in protocol DOCX."""
from __future__ import annotations

import io
from datetime import date

from docx import Document

from services.protocol_docx import WORKSHEET_FORMULAS, fill_protocol_docx_bytes
from services.protocol_store import ProtocolHeader, TestResultRow
from services.protocols.test_catalog import TEST_CATALOG
from services.samples import SampleRecord


def main() -> None:
    sample = SampleRecord(
        id=1,
        request_id=1,
        sample_code="SLS-260717-0001",
        sr_no=1,
        sample_name="Jaggery",
        batch_code="B1",
        quantity="500g",
        parameters="",
        tests_to_perform="",
        status="in_progress",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        lab_code="SLS/CTR/26/001",
        tests_json=(
            '["moisture","total_ash","invert_sugar",'
            '"reducing_sugar","sulphur_dioxide"]'
        ),
    )
    header = ProtocolHeader(
        sample_id=1,
        protocol_no="P-1",
        issued_to="A",
        issued_by="B",
        sample_received_on=date.today(),
        date_of_analysis=date.today(),
        appearance_text="Brown",
    )
    results = [
        TestResultRow(
            "moisture",
            "Moisture",
            "m",
            "%",
            {"empty_dish": 10.0, "w1": 15.0, "w": 5.0, "after_dry": 14.74, "w2": 14.74},
            "5.2",
            5.2,
        ),
        TestResultRow(
            "total_ash",
            "Total ash",
            "m",
            "%",
            {"w1": 20.0, "before_ign": 25.0, "w": 5.0, "after_ign": 20.155, "w2": 20.155},
            "3.1",
            3.1,
        ),
        TestResultRow(
            "invert_sugar",
            "Invert",
            "m",
            "%",
            {
                "sample_wt": 5.0,
                "fehling_invert": 5.0,
                "br_invert": 10.0,
                "sugar_conc": 0.005,
            },
            "92",
            92.0,
        ),
        TestResultRow(
            "reducing_sugar",
            "Reducing",
            "m",
            "%",
            {
                "sample_wt": 5.0,
                "fehling_reducing": 5.0,
                "br_reducing": 12.0,
                "sugar_conc": 0.005,
            },
            "15",
            15.0,
        ),
        TestResultRow(
            "sulphur_dioxide",
            "SO2",
            "m",
            "ppm",
            {"sample_wt": 5.0, "ug_so4": 5.0},
            "10",
            10.0,
        ),
    ]
    raw = fill_protocol_docx_bytes(
        sample,
        header,
        results,
        generated_by="Tester",
        generated_at="2026-07-17 21:00",
    )
    doc = Document(io.BytesIO(raw))
    blob = "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )

    checks = [
        TEST_CATALOG["moisture"].formula_display,
        "Ash w/w (%) = (W2",
        "Total Invert sugar(On dry basis)%",
        "Total reducing sugar %",
        "Sucrose(On dry basis)%",
        "X 250 X 100",
        "X 250 X 10",
        "0.95",
        TEST_CATALOG["sulphur_dioxide"].formula_display,
        # Readings
        "15",  # moisture W1 reading appears
        # Worked moisture: (15 - 14.74) × 100 / 5 = 5.2
        "(15 - 14.74)",
        "5.2",
        # Worked SO2
        "(5 × 10) / 5",
    ]
    failed = 0
    for c in checks:
        ok = c in blob or c.split(";")[0].strip() in blob
        safe = c[:70].encode("ascii", "replace").decode("ascii")
        print(("OK" if ok else "MISSING"), repr(safe))
        if not ok:
            failed += 1

    # Moisture readings column
    moisture_w1 = doc.tables[5].rows[2].cells[1].text.strip()
    print("moisture W1 reading:", repr(moisture_w1))
    if "15" not in moisture_w1:
        print("MISSING moisture reading W1")
        failed += 1

    moisture_formula = doc.tables[5].rows[6].cells[0].text
    print(
        "T5 R6 formula:",
        repr(moisture_formula[:120].encode("ascii", "replace").decode("ascii")),
    )
    if "15 - 14.74" not in moisture_formula and "(15 - 14.74)" not in moisture_formula:
        print("MISSING worked moisture line")
        failed += 1

    moisture_ans = doc.tables[5].rows[6].cells[1].text if len(doc.tables[5].rows[6].cells) > 1 else ""
    print("T5 R6 answer:", repr(moisture_ans))
    if "5.2" not in moisture_ans:
        print("MISSING moisture answer cell")
        failed += 1

    sugar_cell = doc.tables[12].rows[6].cells[0]
    t12 = sugar_cell.text
    print(
        "T12 R6:",
        repr(t12[:200].encode("ascii", "replace").decode("ascii")),
    )
    if "Total Invert sugar" not in t12 or "Total reducing sugar" not in t12:
        print("MISSING sugar formulas on T12")
        failed += 1
    if "Invert answer:" not in t12:
        print("MISSING invert answer on T12")
        failed += 1
    if "Reducing answer:" not in t12:
        print("MISSING reducing answer on T12")
        failed += 1
    if "Sucrose" not in t12 or "0.95" not in t12:
        print("MISSING sucrose formula on T12")
        failed += 1
    # Reducing must use ×10 (Word protocol), not ×100
    paras = sugar_cell.paragraphs
    reducing_para = paras[7].text if len(paras) > 7 else ""
    if "X 250 X 10" not in reducing_para and "× 250 × 10" not in reducing_para:
        # also accept worked line with × 10
        if "× 250 × 10" not in t12 and "X 250 X 10" not in t12:
            print("MISSING reducing ×10 factor on T12")
            failed += 1

    print("tables", len(doc.tables), "map", list(WORKSHEET_FORMULAS))
    for label, text in (
        ("T5 R6", doc.tables[5].rows[6].cells[0].text[:90]),
        ("T6 R6", doc.tables[6].rows[6].cells[0].text[:90]),
        ("T6 R7", doc.tables[6].rows[7].cells[0].text[:90]),
        ("T12 R6", doc.tables[12].rows[6].cells[0].text[:140]),
        ("T15 R2", doc.tables[15].rows[2].cells[0].text[:90]),
    ):
        print(label + ":", repr(text.encode("ascii", "replace").decode("ascii")))
    if failed:
        raise SystemExit(f"{failed} formula check(s) failed")
    print("smoke ok")


if __name__ == "__main__":
    main()
