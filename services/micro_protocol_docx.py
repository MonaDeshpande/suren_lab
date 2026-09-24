"""
services/micro_protocol_docx.py
-------------------------------
Fill reference/Micro Protocol.docx for Micro category analyst protocol.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_LINE_SPACING

from services.document_templates import MICRO_PROTOCOL_PATH
from services.docx_layout import finalize_docx_document
from services.catalog_specs import get_spec
from services.micro_report_catalog import spec_for_key
from services.protocol_docx import (
    _apply_table_borders,
    _apply_table_data_fonts,
    _apply_table_header_row_font,
    _remove_existing_disclaimer_blocks,
    append_protocol_disclaimer,
)
from services.protocol_store import ProtocolHeader, TestResultRow
from services.protocols.test_catalog import MICRO_TEST_KEYS
from services.samples import SampleRecord

MICRO_PROTOCOL_TEMPLATE = MICRO_PROTOCOL_PATH

_MICRO_HEADER_FONT = ("Cambria", 11)
_MICRO_BODY_FONT = ("Cambria", 10)
_MICRO_LINE_SPACING = 1.15


def suggest_micro_protocol_filename(sample: SampleRecord) -> str:
    code = (sample.sample_code or "sample").replace("/", "-")
    return f"Micro_Protocol_{code}.docx"


def _fmt_date(d: Optional[date]) -> str:
    if d is None:
        return ""
    return d.strftime("%d/%m/%Y")


def _set_cell(cell, text: str) -> None:
    """Replace cell content; use cell.text so template runs do not leave stray digits."""
    cell.text = text or ""


def _apply_run_font(run, name: str, size_pt: float, *, bold: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in (qn("w:ascii"), qn("w:hAnsi"), qn("w:cs")):
        r_fonts.set(attr, name)
    half = str(int(round(size_pt * 2)))
    for tag in ("w:sz", "w:szCs"):
        el = r_pr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            r_pr.append(el)
        el.set(qn("w:val"), half)
    b_el = r_pr.find(qn("w:b"))
    if bold:
        if b_el is None:
            r_pr.append(OxmlElement("w:b"))
    elif b_el is not None:
        r_pr.remove(b_el)


def _apply_paragraph_typography(paragraph, *, bold: bool = False) -> None:
    name, size = _MICRO_BODY_FONT if not bold else _MICRO_HEADER_FONT
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = _MICRO_LINE_SPACING
    fmt.space_before = 0
    fmt.space_after = 6
    if not paragraph.runs:
        paragraph.add_run(paragraph.text or "")
    for run in paragraph.runs:
        _apply_run_font(run, name, size, bold=bold)


def _normalize_micro_protocol_layout(doc: Document) -> None:
    """Uniform typography, spacing, and table borders for micro protocols."""
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        is_title = text.lower().startswith(
            ("result", "observation", "protocol", "sample")
        ) or text.endswith(":")
        _apply_paragraph_typography(para, bold=is_title and len(text) < 80)

    for table in doc.tables:
        _apply_table_header_row_font(table)
        _apply_table_data_fonts(table)
        _apply_table_borders(table)


def _result_display(result: Optional[TestResultRow]) -> str:
    if result is None:
        return ""
    value = (result.result_value or "").strip()
    unit = (result.unit or "").strip()
    if value and unit and unit.lower() not in value.lower():
        return f"{value} {unit}".strip()
    return value


def fill_micro_protocol_docx_bytes(
    sample: SampleRecord,
    header: ProtocolHeader,
    results: list[TestResultRow],
    *,
    generated_by: str = "",
    generated_at: str = "",
) -> bytes:
    """Produce a filled Micro protocol .docx from the reference template."""
    if not MICRO_PROTOCOL_TEMPLATE.exists():
        raise FileNotFoundError(
            f"Micro protocol template missing: {MICRO_PROTOCOL_TEMPLATE}"
        )

    doc = Document(io.BytesIO(MICRO_PROTOCOL_TEMPLATE.read_bytes()))
    tables = doc.tables
    if len(tables) < 3:
        raise ValueError("Micro protocol template must have at least 3 tables.")

    by_key = {r.test_key: r for r in results}

    # Table 0 — Protocol No / Issued to / Issued by
    t0 = tables[0]
    if t0.rows and len(t0.rows[0].cells) >= 6:
        row = t0.rows[0]
        _set_cell(row.cells[1], header.protocol_no or "")
        _set_cell(row.cells[3], header.issued_to or "")
        _set_cell(row.cells[5], header.issued_by or "")

    # Table 1 — Sample info
    t1 = tables[1]
    if len(t1.rows) >= 2 and len(t1.columns) >= 4:
        _set_cell(t1.rows[0].cells[1], sample.sample_name or "")
        _set_cell(
            t1.rows[0].cells[3],
            _fmt_date(header.sample_received_on or sample.created_at),
        )
        _set_cell(t1.rows[1].cells[1], sample.sample_code or "")
        _set_cell(
            t1.rows[1].cells[3],
            _fmt_date(header.date_of_analysis),
        )

    # Table 2 — Results (rows 2–7 = tests 1–6)
    t2 = tables[2]
    for index, key in enumerate(MICRO_TEST_KEYS):
        row_idx = index + 2
        if row_idx >= len(t2.rows):
            break
        spec = spec_for_key(key)
        db_spec = get_spec(key)
        name = db_spec.test_name if db_spec else spec.name
        saved = by_key.get(key)
        if saved and (saved.method or "").strip():
            method = saved.method.strip()
        elif db_spec and db_spec.method_of_analysis:
            method = db_spec.method_of_analysis
        else:
            method = spec.method
        row = t2.rows[row_idx]
        _set_cell(row.cells[0], str(index + 1))
        _set_cell(row.cells[1], name)
        _set_cell(row.cells[2], _result_display(by_key.get(key)))
        _set_cell(row.cells[3], method)

    _normalize_micro_protocol_layout(doc)
    _remove_existing_disclaimer_blocks(doc)
    append_protocol_disclaimer(doc, header)
    finalize_docx_document(doc, set_qsf=False)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
