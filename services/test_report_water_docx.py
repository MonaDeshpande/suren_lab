"""
services/test_report_water_docx.py
----------------------------------
Fill reference/Water test report.docx for water sample final reports.

Layout matches the reference multi-section Word document (Chemical + Elemental,
Physical, Microbiological placeholder). PDF conversion reuses docx2pdf when available.
"""

from __future__ import annotations

import io
import json
import logging
import time
from copy import deepcopy
from dataclasses import dataclass, field, replace
from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Length

from services.document_templates import (
    WATER_REPORT_TEMPLATE_PATH,
    final_report_template_path,
)
from services.docx_layout import (
    combine_docx_bytes,
    fill_report_signature_block,
    finalize_docx_document,
    load_template,
    set_cell_text,
    set_paragraph_text,
    set_report_footer,
)
from services.docx_to_pdf import try_convert_docx_to_pdf
from services.protocol_store import ProtocolHeader, TestResultRow
from services.samples import (
    REPORT_FORMAT_BOTH,
    SampleRecord,
    default_report_with_logo,
    default_test_report_no,
    logo_test_keys,
    no_logo_test_keys,
    normalize_report_format,
    report_page_label,
    water_report_keys,
)
from services.water_report_catalog import (
    DEFAULT_TESTING_CONDUCTED_AT,
    WATER_MICRO_PLACEHOLDERS,
    WATER_MICRO_PROTOCOL_KEY_MAP,
    WATER_REPORT_CHEMICAL_TABLE_ROWS,
    WATER_REPORT_ELEMENTAL_TABLE_ROWS,
    WATER_REPORT_LIMITS,
    WATER_REPORT_METHOD_LABELS,
    WATER_REPORT_PHYSICAL_TABLE_ROWS,
    WATER_REPORT_REMARK_CHEMICAL,
    WATER_REPORT_REMARK_MICRO,
    WaterReportLimits,
    limits_for_key,
    method_for_key,
)

logger = logging.getLogger(__name__)

DEFAULT_CUSTOMER_SAMPLE_ID = "Drinking Water"
BLANK_FIELD = "---"
_DEBUG_LOG_PATH = Path(__file__).resolve().parent.parent / "debug-3467ee.log"


def _debug_log(*, hypothesis_id: str, location: str, message: str, data: dict) -> None:
    # region agent log
    try:
        payload = {
            "sessionId": "3467ee",
            "hypothesisId": hypothesis_id,
            "location": location,
            "message": message,
            "data": data,
            "timestamp": int(time.time() * 1000),
        }
        with _DEBUG_LOG_PATH.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(payload) + "\n")
    except OSError:
        pass
    # endregion


def _length_twips(length: Length | int) -> str:
    """OOXML pgMar attributes use twips, not EMU."""
    if isinstance(length, Length):
        return str(int(length.twips))
    return str(int(length))


@dataclass
class WaterReportFillOptions:
    """Reviewer-editable fields for water final report."""

    ulr_no: str = ""
    report_no_chemical: str = ""
    report_no_micro: str = ""
    report_date: Optional[date] = None
    condition_of_sample: str = ""
    customer_sample_id: str = ""
    date_of_sampling: str = BLANK_FIELD
    location_of_sampling: str = BLANK_FIELD
    sampling_method: str = BLANK_FIELD
    sampling_done_by: str = BLANK_FIELD
    sample_appearance: str = ""
    testing_conducted_at: str = DEFAULT_TESTING_CONDUCTED_AT
    test_performance_date: str = ""
    limit_overrides: dict[str, WaterReportLimits] = field(default_factory=dict)
    generated_by: str = ""
    generated_at: str = ""
    authorized_signatory: str = ""
    checked_by: str = ""
    authorized_signatory_role: str = ""
    checked_by_role: str = ""


def _fmt_date(d: Optional[date]) -> str:
    if d is None:
        return ""
    return d.strftime("%d/%m/%Y")


def _fill_signatures(doc: Document, opts: WaterReportFillOptions) -> None:
    fill_report_signature_block(
        doc,
        authorized_name=opts.authorized_signatory,
        authorized_role=opts.authorized_signatory_role,
        checked_name=opts.checked_by,
        checked_role=opts.checked_by_role,
    )


def _result_display(res: TestResultRow) -> str:
    from services.number_format import format_report_number

    val = (res.result_value or "").strip()
    if not val:
        return ""
    formatted = format_report_number(val)
    unit = (res.unit or "").strip()
    if unit and unit not in formatted and unit != "—":
        return f"{formatted} {unit}"
    return formatted


def _default_report_numbers(
    sample: SampleRecord, *, with_logo: bool | None = None
) -> tuple[str, str]:
    logo_flag = default_report_with_logo(sample) if with_logo is None else with_logo
    number = default_test_report_no(sample, with_logo=logo_flag)
    return number, number


def _fill_customer_info_table(
    table,
    sample: SampleRecord,
    header: ProtocolHeader,
    opts: WaterReportFillOptions,
) -> None:
    if len(table.rows) < 9:
        return

    name_addr_parts = [sample.customer_name or ""]
    if (sample.customer_address or "").strip():
        name_addr_parts.append(sample.customer_address.strip())
    name_address = "\n".join(p for p in name_addr_parts if p)

    customer_sample_id = (
        opts.customer_sample_id
        or (sample.parameters or "").strip()
        or DEFAULT_CUSTOMER_SAMPLE_ID
    )
    appearance = opts.sample_appearance or (header.appearance_text or "").strip()
    perf_date = opts.test_performance_date or _fmt_date(header.date_of_analysis)
    testing_at = opts.testing_conducted_at or DEFAULT_TESTING_CONDUCTED_AT

    set_cell_text(table.rows[0].cells[1], name_address)
    set_cell_text(table.rows[1].cells[1], customer_sample_id)
    set_cell_text(table.rows[1].cells[3], sample.batch_code or BLANK_FIELD)
    set_cell_text(table.rows[2].cells[1], sample.lab_code or sample.sample_code or "")
    set_cell_text(table.rows[2].cells[3], opts.date_of_sampling or BLANK_FIELD)
    set_cell_text(table.rows[3].cells[1], _fmt_date(header.sample_received_on))
    set_cell_text(table.rows[3].cells[3], opts.location_of_sampling or BLANK_FIELD)
    set_cell_text(table.rows[4].cells[1], sample.sample_name or "")
    set_cell_text(table.rows[5].cells[1], opts.sampling_method or BLANK_FIELD)
    set_cell_text(table.rows[5].cells[3], opts.sampling_done_by or BLANK_FIELD)
    set_cell_text(
        table.rows[6].cells[1],
        opts.condition_of_sample or "",
    )
    set_cell_text(table.rows[6].cells[3], perf_date)
    set_cell_text(table.rows[7].cells[1], appearance)
    set_cell_text(table.rows[7].cells[3], sample.quantity or "")
    set_cell_text(table.rows[8].cells[1], testing_at)


def _fill_result_row(
    table,
    row_idx: int,
    test_key: str,
    by_key: dict[str, TestResultRow],
    limit_overrides: dict[str, WaterReportLimits],
) -> None:
    if row_idx >= len(table.rows):
        return
    row = table.rows[row_idx]
    if len(row.cells) < 6:
        return
    limits = limits_for_key(test_key, limit_overrides)
    res = by_key.get(test_key)
    result_text = _result_display(res) if res else ""
    set_cell_text(row.cells[2], result_text)
    if limits.desirable:
        set_cell_text(row.cells[3], limits.desirable)
    if limits.permissible:
        set_cell_text(row.cells[4], limits.permissible)
    method = method_for_key(test_key)
    if method:
        set_cell_text(row.cells[5], method)


def _fill_chemical_elemental_table(
    table,
    by_key: dict[str, TestResultRow],
    limit_overrides: dict[str, WaterReportLimits],
) -> None:
    for row_idx, test_key in WATER_REPORT_CHEMICAL_TABLE_ROWS.items():
        _fill_result_row(table, row_idx, test_key, by_key, limit_overrides)
    for row_idx, test_key in WATER_REPORT_ELEMENTAL_TABLE_ROWS.items():
        _fill_result_row(table, row_idx, test_key, by_key, limit_overrides)
    if len(table.rows) > 12:
        set_cell_text(table.rows[12].cells[0], WATER_REPORT_REMARK_CHEMICAL)


def _fill_physical_table(
    table,
    by_key: dict[str, TestResultRow],
    limit_overrides: dict[str, WaterReportLimits],
) -> None:
    for row_idx, test_key in WATER_REPORT_PHYSICAL_TABLE_ROWS.items():
        _fill_result_row(table, row_idx, test_key, by_key, limit_overrides)


def _micro_result_display(res: TestResultRow | None) -> str:
    """Present/Absent from water micro protocol observation."""
    if res is None:
        return ""
    val = (res.result_value or "").strip()
    if not val:
        val = str((res.inputs or {}).get("result_obs") or "").strip()
    if not val:
        return ""
    lower = val.lower()
    if lower == "present":
        return "Present"
    if lower == "absent":
        return "Absent"
    return val


def _find_page2_start_index(doc: Document) -> int | None:
    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip().upper()
        if "7.8" in text and ("QSF" in text or "QSP" in text):
            return i
    return None


def _clear_blank_paragraphs(doc: Document, indices: set[int]) -> None:
    for i in indices:
        if i < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[i], "")


def _insert_page2_section_break(doc: Document, break_before_index: int) -> None:
    """End section 1 and start section 2 with page numbering restarted at 1."""
    if break_before_index <= 0 or break_before_index >= len(doc.paragraphs):
        return
    break_para = doc.paragraphs[break_before_index - 1]
    p_pr = break_para._element.get_or_add_pPr()
    for existing in list(p_pr.findall(qn("w:sectPr"))):
        p_pr.remove(existing)

    first_section = doc.sections[0]
    sect_pr = OxmlElement("w:sectPr")
    sect_type = OxmlElement("w:type")
    sect_type.set(qn("w:val"), "nextPage")
    sect_pr.append(sect_type)

    pg_sz = first_section._sectPr.find(qn("w:pgSz"))
    if pg_sz is not None:
        sect_pr.append(deepcopy(pg_sz))

    pg_mar = OxmlElement("w:pgMar")
    pg_mar.set(qn("w:top"), _length_twips(Inches(0.15)))
    pg_mar.set(qn("w:right"), _length_twips(first_section.right_margin))
    pg_mar.set(qn("w:bottom"), _length_twips(first_section.bottom_margin))
    pg_mar.set(qn("w:left"), _length_twips(first_section.left_margin))
    pg_mar.set(qn("w:header"), _length_twips(Inches(0.35)))
    pg_mar.set(qn("w:footer"), _length_twips(Inches(0.35)))
    sect_pr.append(pg_mar)

    _debug_log(
        hypothesis_id="A",
        location="test_report_water_docx.py:_insert_page2_section_break",
        message="inline sectPr pgMar twips",
        data={
            "top": pg_mar.get(qn("w:top")),
            "right": pg_mar.get(qn("w:right")),
            "header": pg_mar.get(qn("w:header")),
            "footer": pg_mar.get(qn("w:footer")),
        },
    )

    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    sect_pr.append(pg_num)

    header_ref = first_section._sectPr.find(qn("w:headerReference"))
    if header_ref is not None:
        sect_pr.append(deepcopy(header_ref))
    footer_ref = first_section._sectPr.find(qn("w:footerReference"))
    if footer_ref is not None:
        sect_pr.append(deepcopy(footer_ref))

    p_pr.append(sect_pr)

    if len(doc.sections) > 1:
        doc.sections[1].footer.is_linked_to_previous = False


def _configure_water_report_sections(doc: Document) -> None:
    """Split into two single-page sections; tighten page-2 header spacing."""
    page2_start = _find_page2_start_index(doc)
    if page2_start is None:
        return

    spacer_indices = {
        page2_start + 1,
        page2_start + 4,
        page2_start + 5,
        page2_start + 9,
        page2_start + 10,
    }
    _clear_blank_paragraphs(doc, spacer_indices)
    _insert_page2_section_break(doc, page2_start)


def _fill_micro_table(table, by_key: dict[str, TestResultRow]) -> None:
    """Fill microbiological rows from water protocol observation results."""
    if len(table.rows) < 4:
        return
    for i, placeholder in enumerate(WATER_MICRO_PLACEHOLDERS):
        row_idx = 2 + i
        if row_idx >= len(table.rows):
            break
        row = table.rows[row_idx]
        if len(row.cells) < 6:
            continue
        protocol_key = WATER_MICRO_PROTOCOL_KEY_MAP.get(placeholder.key, placeholder.key)
        res = by_key.get(protocol_key)
        set_cell_text(row.cells[0], placeholder.sr_no)
        set_cell_text(row.cells[1], placeholder.name)
        set_cell_text(row.cells[2], _micro_result_display(res))
        set_cell_text(row.cells[3], placeholder.desirable)
        if placeholder.permissible:
            set_cell_text(row.cells[4], placeholder.permissible)
        set_cell_text(row.cells[5], placeholder.method)
    if len(table.rows) > 4:
        set_cell_text(table.rows[4].cells[0], WATER_REPORT_REMARK_MICRO)


def _fill_header_paragraphs(
    doc: Document, opts: WaterReportFillOptions, *, with_logo: bool
) -> None:
    report_date = _fmt_date(opts.report_date) if opts.report_date else ""
    ulr = (opts.ulr_no or "").strip()
    chem_report = (opts.report_no_chemical or "").strip()
    micro_report = (opts.report_no_micro or "").strip()

    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if i == 0 and text.startswith("ULR No:"):
            if with_logo:
                set_paragraph_text(
                    para, f"ULR No: {ulr}" if ulr else "ULR No:"
                )
            else:
                set_paragraph_text(para, "")
        elif "Page 1 of 1" in text or "page 1 of 1" in text.lower():
            set_paragraph_text(para, report_page_label(with_logo=with_logo))
        elif i == 2 and text.startswith("Date:"):
            set_paragraph_text(para, f"Date: {report_date}" if report_date else "Date:")
        elif i == 3 and "Report No:" in text:
            set_paragraph_text(
                para,
                f"Report No:  {chem_report}" if chem_report else "Report No:",
            )
        elif i == 29 and "Report No:" in text:
            date_part = f"Date: {report_date}" if report_date else "Date:"
            report_part = f"Report No: {micro_report}" if micro_report else "Report No:"
            set_paragraph_text(para, f"{date_part}\t\t\t\t                       {report_part}")


def fill_water_test_report_docx_bytes(
    sample: SampleRecord,
    header: ProtocolHeader,
    results: list[TestResultRow],
    opts: Optional[WaterReportFillOptions] = None,
    *,
    test_key_filter: Optional[set[str]] = None,
    with_logo: bool | None = None,
) -> bytes:
    """Produce a filled water test report .docx."""
    fill_opts = opts or WaterReportFillOptions()
    by_key = {r.test_key: r for r in results}
    full_by_key = dict(by_key)
    if test_key_filter is not None:
        by_key = {k: v for k, v in by_key.items() if k in test_key_filter}
    logo_flag = default_report_with_logo(sample) if with_logo is None else with_logo
    template_path = final_report_template_path(sample.category, with_logo=logo_flag)
    if not template_path.exists():
        raise FileNotFoundError(
            f"Water test report template missing: {template_path}"
        )

    if not fill_opts.report_no_chemical or not fill_opts.report_no_micro:
        default_chem, default_micro = _default_report_numbers(
            sample, with_logo=logo_flag
        )
        if not fill_opts.report_no_chemical:
            fill_opts.report_no_chemical = default_chem
        if not fill_opts.report_no_micro:
            fill_opts.report_no_micro = default_micro

    if not fill_opts.report_date:
        fill_opts.report_date = header.date_of_analysis or date.today()

    if not fill_opts.sample_appearance:
        fill_opts.sample_appearance = (header.appearance_text or "").strip()

    doc = load_template(template_path)
    _fill_signatures(doc, fill_opts)
    _fill_header_paragraphs(doc, fill_opts, with_logo=logo_flag)

    tables = doc.tables
    if len(tables) >= 1:
        _fill_customer_info_table(tables[0], sample, header, fill_opts)
    if len(tables) >= 2:
        _fill_chemical_elemental_table(
            tables[1], by_key, fill_opts.limit_overrides
        )
    if len(tables) >= 3:
        _fill_customer_info_table(tables[2], sample, header, fill_opts)
    if len(tables) >= 4:
        _fill_physical_table(tables[3], by_key, fill_opts.limit_overrides)
    if len(tables) >= 5:
        _fill_micro_table(tables[4], full_by_key)

    _configure_water_report_sections(doc)
    set_report_footer(doc, with_logo=logo_flag, static_one_of_one=True)

    finalize_docx_document(doc, set_qsf=True)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def suggest_water_test_report_filename(
    sample: SampleRecord,
    *,
    extension: str = "docx",
) -> str:
    safe = "".join(
        ch if ch.isalnum() or ch in "-_" else "_"
        for ch in (sample.sample_name or sample.sample_code or "sample")
    )
    ext = extension.lstrip(".")
    return f"TestReport_{safe}_{sample.sample_code}.{ext}"


def generate_water_test_report_pdf_bytes(
    sample: SampleRecord,
    header: ProtocolHeader,
    results: list[TestResultRow],
    opts: Optional[WaterReportFillOptions] = None,
) -> tuple[bytes, bytes | None]:
    """
    Returns (docx_bytes, pdf_bytes_or_none).
    PDF is None when Word / LibreOffice conversion is unavailable.
    """
    fmt = normalize_report_format(sample.report_format)
    if fmt == REPORT_FORMAT_BOTH:
        parts: list[bytes] = []
        logo_flags: list[bool] = []
        base_opts = opts or WaterReportFillOptions()
        if logo_test_keys(sample):
            num = (
                (base_opts.report_no_chemical or "").strip()
                or default_test_report_no(sample, with_logo=True)
            )
            logo_opts = replace(
                base_opts, report_no_chemical=num, report_no_micro=num
            )
            parts.append(
                fill_water_test_report_docx_bytes(
                    sample,
                    header,
                    results,
                    opts=logo_opts,
                    test_key_filter=logo_test_keys(sample),
                    with_logo=True,
                )
            )
            logo_flags.append(True)
        if no_logo_test_keys(sample):
            num = default_test_report_no(sample, with_logo=False)
            nologo_opts = replace(
                base_opts, report_no_chemical=num, report_no_micro=num
            )
            parts.append(
                fill_water_test_report_docx_bytes(
                    sample,
                    header,
                    results,
                    opts=nologo_opts,
                    test_key_filter=no_logo_test_keys(sample),
                    with_logo=False,
                )
            )
            logo_flags.append(False)
        docx_bytes = (
            combine_docx_bytes(
                parts,
                with_logo_per_section=logo_flags,
                static_one_of_one=True,
            )
            if parts
            else fill_water_test_report_docx_bytes(sample, header, results, opts=opts)
        )
    else:
        docx_bytes = fill_water_test_report_docx_bytes(
            sample, header, results, opts=opts
        )
    pdf_bytes = try_convert_docx_to_pdf(docx_bytes)
    return docx_bytes, pdf_bytes


def build_water_report_limit_rows(
    sample: SampleRecord,
    results: list[TestResultRow],
    limit_overrides: Optional[dict[str, WaterReportLimits]] = None,
) -> list[dict[str, str]]:
    """Preview rows for Reviewer spec editor (Desirable / Permissible columns)."""
    by_key = {r.test_key: r for r in results}
    overrides = limit_overrides or {}
    rows: list[dict[str, str]] = []
    for key in water_report_keys(sample):
        limits = limits_for_key(key, overrides)
        res = by_key.get(key)
        result_val = _result_display(res) if res else ""
        rows.append(
            {
                "Test": key,
                "Result": result_val,
                "Desirable Limit": limits.desirable,
                "Permissible limit": limits.permissible,
                "Method": method_for_key(key),
            }
        )
    return rows
