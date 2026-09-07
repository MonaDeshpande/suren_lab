"""
services/docx_layout.py
-----------------------
Shared helpers for filling Word templates while preserving reference layout
(fonts, paragraph alignment, row heights, section margins).
"""

from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Row
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

from docx.shared import Pt

from services.samples import report_page_label

REPORT_QSF_LABEL = "QSF No -7.8.2"
AUTHORIZED_SIGNATORY_MARKER = "{{AUTHORIZED_SIGNATORY}}"
CHECKED_BY_MARKER = "{{CHECKED_BY}}"
LAB_SHORT_NAME = "SLS"

_SUBSCRIPT_CHARS = "₀₁₂₃₄₅₆₇₈₉"
_SUPERSCRIPT_CHARS = "⁰¹²³⁴⁵⁶⁷⁸⁹"
_EXTRA_SUPERSCRIPT = {"²": "2", "³": "3"}
_SUBSCRIPT_MAP = str.maketrans(_SUBSCRIPT_CHARS, "0123456789")
_SUPERSCRIPT_MAP = str.maketrans(_SUPERSCRIPT_CHARS, "0123456789")
_SCRIPT_CHAR_SET = frozenset(_SUBSCRIPT_CHARS + _SUPERSCRIPT_CHARS + "²³")


def load_template(path: Path) -> Document:
    """Open a reference .docx template (bytes copy — safe to mutate)."""
    if not path.exists():
        raise FileNotFoundError(f"Template missing: {path}")
    return Document(io.BytesIO(path.read_bytes()))


def set_paragraph_text(paragraph, text: str) -> None:
    """Replace paragraph text, preserving first-run font properties."""
    text = text or ""
    if not paragraph.runs:
        paragraph.text = text
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Replace cell text in-place, preserving template paragraph formatting."""
    text = text or ""
    if not cell.paragraphs:
        cell.text = text
        return
    set_paragraph_text(cell.paragraphs[0], text)
    for para in cell.paragraphs[1:]:
        set_paragraph_text(para, "")


def find_paragraph(doc: Document, contains: str) -> int | None:
    """Return paragraph index whose text contains *contains*, else None."""
    needle = (contains or "").strip().lower()
    for i, para in enumerate(doc.paragraphs):
        if needle in (para.text or "").lower():
            return i
    return None


def set_paragraph_by_marker(doc: Document, marker: str, new_text: str) -> bool:
    """Replace entire paragraph text when it contains *marker*."""
    idx = find_paragraph(doc, marker)
    if idx is None:
        return False
    set_paragraph_text(doc.paragraphs[idx], new_text)
    return True


def _iter_all_paragraphs(doc: Document):
    """Yield every paragraph in body, tables, headers, and footers."""
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para
    for section in doc.sections:
        for para in section.header.paragraphs:
            yield para
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        yield para
        for para in section.footer.paragraphs:
            yield para
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        yield para


def set_report_qsf_label(doc: Document, label: str = REPORT_QSF_LABEL) -> None:
    """Normalize QSF/QSP header text on customer final reports."""
    for para in _iter_all_paragraphs(doc):
        text = (para.text or "").strip()
        upper = text.upper()
        if "7.8" in text and ("QSF" in upper or "QSP" in upper):
            set_paragraph_text(para, label)


def _split_script_segments(text: str) -> list[tuple[str, str | None]]:
    """Split text into (segment, script) where script is sub/super/None."""
    segments: list[tuple[str, str | None]] = []
    for ch in text:
        if ch in _SUBSCRIPT_CHARS:
            digit = ch.translate(_SUBSCRIPT_MAP)
            if segments and segments[-1][1] == "sub":
                prev, _ = segments[-1]
                segments[-1] = (prev + digit, "sub")
            else:
                segments.append((digit, "sub"))
        elif ch in _SUPERSCRIPT_CHARS:
            digit = ch.translate(_SUPERSCRIPT_MAP)
            if segments and segments[-1][1] == "super":
                prev, _ = segments[-1]
                segments[-1] = (prev + digit, "super")
            else:
                segments.append((digit, "super"))
        elif ch in _EXTRA_SUPERSCRIPT:
            digit = _EXTRA_SUPERSCRIPT[ch]
            if segments and segments[-1][1] == "super":
                prev, _ = segments[-1]
                segments[-1] = (prev + digit, "super")
            else:
                segments.append((digit, "super"))
        else:
            if segments and segments[-1][1] is None:
                prev, _ = segments[-1]
                segments[-1] = (prev + ch, None)
            else:
                segments.append((ch, None))
    return segments


def _rewrite_paragraph_scripts(paragraph) -> None:
    text = paragraph.text or ""
    if not any(ch in _SCRIPT_CHAR_SET for ch in text):
        return

    base = paragraph.runs[0] if paragraph.runs else None
    font_name = base.font.name if base is not None else None
    font_size = base.font.size if base is not None else None
    bold = base.bold if base is not None else None
    italic = base.italic if base is not None else None

    p_element = paragraph._p
    for run_el in list(p_element.findall(qn("w:r"))):
        p_element.remove(run_el)

    for segment_text, script in _split_script_segments(text):
        if not segment_text:
            continue
        run = paragraph.add_run(segment_text)
        if font_name:
            run.font.name = font_name
        if font_size is not None:
            run.font.size = font_size
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if script == "sub":
            run.font.subscript = True
        elif script == "super":
            run.font.superscript = True


def rewrite_unicode_scripts_in_document(doc: Document) -> None:
    """Replace Unicode sub/superscript glyphs with Word vertAlign runs for PDF export."""
    for para in _iter_all_paragraphs(doc):
        _rewrite_paragraph_scripts(para)


def rewrite_unicode_scripts_in_docx_bytes(docx_bytes: bytes) -> bytes:
    """Rewrite Unicode scripts in a .docx blob before PDF conversion."""
    if not docx_bytes or docx_bytes[:2] != b"PK":
        return docx_bytes
    try:
        doc = Document(io.BytesIO(docx_bytes))
    except Exception:
        return docx_bytes
    rewrite_unicode_scripts_in_document(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def finalize_docx_document(doc: Document, *, set_qsf: bool = False) -> None:
    """Apply pre-save fixes shared by report/protocol fillers."""
    rewrite_unicode_scripts_in_document(doc)
    if set_qsf:
        set_report_qsf_label(doc)


def duplicate_table_row(table: Table, template_row_index: int) -> _Row:
    """Append a copy of an existing row (preserves trPr / cell widths)."""
    template_tr = table.rows[template_row_index]._tr
    new_tr = deepcopy(template_tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def insert_table_row_before(table: Table, before_index: int, template_row_index: int) -> _Row:
    """Insert a row copy before *before_index*."""
    template_tr = table.rows[template_row_index]._tr
    new_tr = deepcopy(template_tr)
    table.rows[before_index]._tr.addprevious(new_tr)
    return table.rows[before_index]


def ensure_table_rows(
    table: Table,
    *,
    first_data_row: int,
    needed_data_rows: int,
    template_data_row: int | None = None,
) -> None:
    """
    Grow *table* so rows [first_data_row .. first_data_row+needed-1] exist.

    Clones *template_data_row* (defaults to first_data_row) for new rows.
    """
    if needed_data_rows <= 0:
        return
    clone_from = template_data_row if template_data_row is not None else first_data_row
    last_needed = first_data_row + needed_data_rows - 1
    while len(table.rows) <= last_needed:
        duplicate_table_row(table, clone_from)


def delete_table_row(table: Table, row_index: int) -> None:
    """Remove a row from *table* by index."""
    tr = table.rows[row_index]._tr
    table._tbl.remove(tr)


def delete_table_rows_range(table: Table, start_index: int, end_index: int) -> None:
    """Remove rows [start_index, end_index) — delete from bottom up."""
    if start_index >= end_index:
        return
    for idx in range(end_index - 1, start_index - 1, -1):
        delete_table_row(table, idx)


def _section_footer_parts(section) -> list:
    """Return every footer variant on a section (default, first, even)."""
    parts = [section.footer]
    if section.different_first_page_header_footer:
        parts.append(section.first_page_footer)
    if section._sectPr.find(qn("w:evenAndOddHeaders")) is not None:
        parts.append(section.even_page_footer)
    return parts


def _clear_footer_part(footer) -> None:
    """Remove all paragraphs, tables, and field remnants from a footer."""
    footer.is_linked_to_previous = False
    el = footer._element
    for child in list(el):
        el.remove(child)


def _append_field_run(paragraph: Paragraph, field_code: str, placeholder: str = "1") -> None:
    """Append a Word field (PAGE, NUMPAGES, etc.) to *paragraph*."""
    run = paragraph.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("w:space"), "preserve")
    instr.text = field_code

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    result_run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = placeholder
    result_run.append(text_el)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    paragraph._p.append(result_run)
    r.append(fld_end)


def _set_footer_page_label(footer, *, with_logo: bool) -> None:
    """Replace footer content with one right-aligned 'page N of M' field line."""
    _clear_footer_part(footer)
    para = footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    word = "page" if with_logo else "pg"
    para.add_run(f"{word} ")
    _append_field_run(para, " PAGE ")
    para.add_run(" of ")
    _append_field_run(para, " NUMPAGES ")


def _set_footer_static_one_of_one(footer, *, with_logo: bool) -> None:
    """Replace footer with a fixed 'page 1 of 1' / 'pg 1 of 1' label."""
    _clear_footer_part(footer)
    para = footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.add_run(report_page_label(with_logo=with_logo))


def compact_single_page_document(doc: Document) -> None:
    """Tighten vertical spacing so typical food/micro reports stay on one page."""
    for para in _iter_all_paragraphs(doc):
        text = (para.text or "").strip()
        pf = para.paragraph_format
        if text in (
            "Disclaimer",
            "End of Report",
            "CHEMICAL TEST REPORT",
            "Microbiological Test",
            "TEST REPORT",
        ) or text.startswith(("Remark:", "Disclaimer")):
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
        else:
            if pf.space_before is not None and pf.space_before.pt > 6:
                pf.space_before = Pt(3)
            if pf.space_after is not None and pf.space_after.pt > 6:
                pf.space_after = Pt(3)


def fill_report_signature_block(
    doc: Document,
    *,
    authorized_name: str = "",
    authorized_role: str = "",
    checked_name: str = "",
    checked_role: str = "",
    lab_short_name: str = LAB_SHORT_NAME,
) -> None:
    """Fill authorized / checked-by blocks on final report templates."""
    auth = (authorized_name or "").strip()
    auth_role = (authorized_role or "Director").strip()
    checked = (checked_name or "").strip()
    checked_role_text = (checked_role or "Quality Manager").strip()
    checked_line = f"Checked by: {checked}" if checked else "Checked by:"

    for para in doc.paragraphs:
        text = para.text or ""
        if AUTHORIZED_SIGNATORY_MARKER in text or CHECKED_BY_MARKER in text:
            new_text = text.replace(AUTHORIZED_SIGNATORY_MARKER, auth)
            new_text = new_text.replace(CHECKED_BY_MARKER, checked)
            if CHECKED_BY_MARKER not in text and "Checked by:" in text:
                new_text = new_text.replace("Checked by:", checked_line, 1)
            set_paragraph_text(para, new_text)
        elif text.startswith("Dr.") or text.startswith("Mrs."):
            set_paragraph_text(para, auth)
        elif text.strip() == "Director" and auth_role:
            set_paragraph_text(para, auth_role)
        elif text.strip() == "Quality Manager" and checked_role_text:
            set_paragraph_text(para, checked_role_text)
        elif text.strip().startswith("Checked by"):
            set_paragraph_text(para, checked_line)


def set_ctr_signature_footer(
    doc: Document,
    *,
    left_text: str,
    right_text: str = "Review",
) -> None:
    """Stamp Reception name/date (left) and Review (right) on every page footer."""
    for section in doc.sections:
        for footer_part in _section_footer_parts(section):
            _clear_footer_part(footer_part)
            left_para = footer_part.add_paragraph()
            left_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            left_para.add_run(left_text or "")
            right_para = footer_part.add_paragraph()
            right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right_para.add_run(right_text)


def set_report_footer(
    doc: Document,
    *,
    with_logo: bool,
    page: int = 1,
    total: int = 1,
    with_logo_per_section: list[bool] | None = None,
    static_one_of_one: bool = False,
) -> None:
    """
    Right-align a single page label in every section footer.

    When *static_one_of_one* is True, writes literal ``page 1 of 1`` text.
    Otherwise uses Word PAGE / NUMPAGES fields (``page`` / *total* ignored).
    """
    for i, section in enumerate(doc.sections):
        logo_flag = (
            with_logo_per_section[i]
            if with_logo_per_section is not None and i < len(with_logo_per_section)
            else with_logo
        )
        for footer_part in _section_footer_parts(section):
            if static_one_of_one:
                _set_footer_static_one_of_one(footer_part, with_logo=logo_flag)
            else:
                _set_footer_page_label(footer_part, with_logo=logo_flag)


def clear_header_images(doc: Document) -> None:
    """Remove drawings from headers (legacy helper — final reports use separate templates)."""
    for section in doc.sections:
        header = section.header
        for para in header.paragraphs:
            for run in para.runs:
                for child in list(run._r):
                    if "drawing" in child.tag or child.tag.endswith("drawing"):
                        run._r.remove(child)
        for tbl in header.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            for child in list(run._r):
                                if "drawing" in child.tag or child.tag.endswith("drawing"):
                                    run._r.remove(child)


def combine_docx_bytes(
    parts: list[bytes],
    *,
    with_logo_per_section: list[bool] | None = None,
    static_one_of_one: bool = False,
) -> bytes:
    """Concatenate filled reports with a new page/section per part."""
    if not parts:
        raise ValueError("No documents to combine")
    if len(parts) == 1:
        return parts[0]
    dst = Document(io.BytesIO(parts[0]))
    for blob in parts[1:]:
        src = Document(io.BytesIO(blob))
        dst.add_section(WD_SECTION.NEW_PAGE)
        dst_body = dst.element.body
        for child in list(src.element.body):
            if child.tag == qn("w:sectPr"):
                continue
            dst_body.insert(-1, deepcopy(child))
    if with_logo_per_section is None:
        with_logo_per_section = [True] * len(dst.sections)
    set_report_footer(
        dst,
        with_logo=True,
        with_logo_per_section=with_logo_per_section,
        static_one_of_one=static_one_of_one,
    )
    out = io.BytesIO()
    dst.save(out)
    return out.getvalue()
