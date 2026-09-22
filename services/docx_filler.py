"""
services/docx_filler.py
-----------------------
Fill the reference Word template with form data and return .docx bytes.

Template:
  reference/Customer Test Request form LLP.docx

PDF download uses ReportLab layout via services/ctr_pdf.py.
This module serves the Word (.docx) download (LLP body + CTR_template header only).
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

from services.ctr_verification import (
    CHECKLIST_TITLE,
    CUSTOMER_SIGNATURE_LABEL,
    SAMPLE_SECTION_HEADING,
    SAMPLE_TABLE_HEADERS,
    ctr_signature_date,
    ctr_test_names_for_sample,
    verification_checklist_rows,
)
from services.ctr_symbols import bracket_mark, yes_no_option_line
from services.customers import format_contacts_for_display
from services.requests import (
    SampleRow,
    TestRequestData,
    ctr_parameters_display,
    delivery_mode_is_selected,
    effective_sample_request_details,
)
from docx.text.paragraph import Paragraph

from services.document_templates import CTR_FORM_BODY_PATH, CTR_LETTERHEAD_PATH
from services.docx_layout import (
    apply_ctr_header_only,
    apply_table_cell_font_size,
    compact_paragraph_spacing,
    find_paragraph,
    read_table_grid_column_widths,
    remove_body_paragraph_containing,
    remove_empty_paragraphs_following_table,
    remove_table,
    replace_table_grid_columns,
)

TEMPLATE_PATH = CTR_FORM_BODY_PATH

# ~6.5" content width in twips (matches ReportLab usable width).
_CTR_CONTENT_WIDTH_TWIPS = 9360
_CTR_CHECKLIST_COL_WIDTHS = [650, 3400, _CTR_CONTENT_WIDTH_TWIPS - 650 - 3400]
_CTR_DEFAULT_SAMPLE_COL_WIDTHS = [720, 2000, 1600, 1200, 3840]
_STALE_SAMPLE_DESCRIPTION_NEEDLE = "Sample Description & tests to be perform"


def _fmt_date(d: Optional[date]) -> str:
    if d is None:
        return ""
    return d.strftime("%d/%m/%Y")


def _yes_no_cell(value: Optional[bool]) -> str:
    """Match the Word template Yes/No option cell style."""
    return yes_no_option_line(value)


def _service_cell(service_type: str) -> str:
    """Urgent / Regular options as on the Word form."""
    s = (service_type or "").strip().lower()
    urgent = bracket_mark(s == "urgent")
    regular = bracket_mark(s == "regular")
    return f"Urgent\t{urgent}\t\t\tRegular\t{regular}"


def _delivery_cell(mode: str) -> str:
    """Collect / Courier / Email/Whatsapp options as on the Word form."""
    collect = bracket_mark(delivery_mode_is_selected(mode, "Collect"))
    courier = bracket_mark(delivery_mode_is_selected(mode, "Courier"))
    email = bracket_mark(delivery_mode_is_selected(mode, "Email/Whatsapp"))
    return f"Collect\t{collect}\t Courier\t{courier}\t  Email/Whatsapp\t{email}"


def _set_cell_text(cell, text: str) -> None:
    """
    Replace all paragraph text inside a table cell with `text`.

    Keeps the first paragraph (preserves some formatting) and clears extras.
    """
    text = text or ""
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.text = text
        return

    paragraphs[0].text = text
    for p in paragraphs[1:]:
        p.text = ""


def _append_to_label(cell, label_prefix: str, value: str) -> None:
    """Write 'Label: value' into a cell (template already has the label)."""
    value = value or ""
    _set_cell_text(cell, f"{label_prefix} {value}".strip())


def _add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def _style_ctr_heading(paragraph: Paragraph) -> None:
    compact_paragraph_spacing(paragraph, space_before_pt=4, space_after_pt=2)


def _add_ctr_bold_heading(doc: Document, text: str) -> Paragraph:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    _style_ctr_heading(para)
    return para


def _set_table_borders(table: Table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def _add_sample_table(
    doc: Document,
    sample: SampleRow,
    sr_index: int,
    *,
    col_widths: list[int] | None = None,
) -> Table:
    table = doc.add_table(rows=2, cols=5)
    _set_table_borders(table)
    widths = col_widths or _CTR_DEFAULT_SAMPLE_COL_WIDTHS
    if len(widths) == 5:
        replace_table_grid_columns(table, widths)
    header_cells = table.rows[0].cells
    for idx, label in enumerate(SAMPLE_TABLE_HEADERS):
        _set_cell_text(header_cells[idx], label)
    data_cells = table.rows[1].cells
    _set_cell_text(data_cells[0], str(sr_index))
    _set_cell_text(data_cells[1], sample.sample_name or "")
    _set_cell_text(data_cells[2], sample.batch_code or "")
    _set_cell_text(data_cells[3], sample.quantity or "")
    _set_cell_text(data_cells[4], ctr_parameters_display(sample))
    apply_table_cell_font_size(table, size_pt=10)
    return table


def _add_tests_paragraphs(doc: Document, sample: SampleRow) -> None:
    names = ctr_test_names_for_sample(sample)
    if not names:
        return
    _add_ctr_bold_heading(doc, "Tests to be performed:")
    for index, name in enumerate(names, start=1):
        line = doc.add_paragraph(f"{index}. {name}")
        compact_paragraph_spacing(line, space_before_pt=0, space_after_pt=0)

def _add_verification_table(
    doc: Document,
    sample: SampleRow,
    *,
    storage_temperature: str = "",
) -> Table:
    rows = verification_checklist_rows(
        sample,
        storage_temperature=storage_temperature,
    )
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    _set_table_borders(table)
    header = table.rows[0].cells
    _set_cell_text(header[0], "Sr No")
    _set_cell_text(header[1], "Particulars")
    _set_cell_text(header[2], "Remark")
    for row_idx, row in enumerate(rows, start=1):
        cells = table.rows[row_idx].cells
        _set_cell_text(cells[0], str(row.sr))
        _set_cell_text(cells[1], row.particular)
        _set_cell_text(cells[2], row.remark)
    replace_table_grid_columns(table, _CTR_CHECKLIST_COL_WIDTHS)
    apply_table_cell_font_size(table, size_pt=10)
    return table


def _apply_ctr_signature_block(
    doc: Document,
    *,
    generated_by: str,
    generated_at: str,
    request_date: Optional[date],
) -> None:
    """Replace template Receiver line with name, date, Receiver | customer sign."""
    idx = find_paragraph(doc, "Sign & date")
    if idx is not None:
        element = doc.paragraphs[idx]._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    else:
        remove_body_paragraph_containing(doc, "Receiver")

    if len(doc.tables) < 2:
        return

    note_table = doc.tables[1]
    remove_empty_paragraphs_following_table(doc, note_table, max_remove=15)

    reception_name = (generated_by or "Reception").strip()
    sig_date = ctr_signature_date(generated_at, request_date)

    sign_table = doc.add_table(rows=3, cols=2)
    replace_table_grid_columns(
        sign_table,
        [_CTR_CONTENT_WIDTH_TWIPS // 2, _CTR_CONTENT_WIDTH_TWIPS // 2],
    )
    _set_cell_text(sign_table.rows[0].cells[0], reception_name)
    _set_cell_text(sign_table.rows[0].cells[1], CUSTOMER_SIGNATURE_LABEL)
    _set_cell_text(sign_table.rows[1].cells[0], sig_date)
    _set_cell_text(sign_table.rows[1].cells[1], "")
    _set_cell_text(sign_table.rows[2].cells[0], "Receiver")
    _set_cell_text(sign_table.rows[2].cells[1], "")
    apply_table_cell_font_size(sign_table, size_pt=11)

    note_tbl = note_table._tbl
    tbl_el = sign_table._tbl
    tbl_el.getparent().remove(tbl_el)

    blank1 = OxmlElement("w:p")
    blank2 = OxmlElement("w:p")
    note_tbl.addnext(blank1)
    blank1.addnext(blank2)
    blank2.addnext(tbl_el)


def fill_docx_bytes(
    data: TestRequestData,
    generated_by: str = "",
    generated_at: str = "",
) -> bytes:
    """
    Open the reference template, fill tables, return .docx as bytes.

    Raises
    ------
    FileNotFoundError
        If the Word template is missing from /reference.
    """
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Word template not found at: {TEMPLATE_PATH}")
    if not CTR_LETTERHEAD_PATH.exists():
        raise FileNotFoundError(f"CTR letterhead template not found at: {CTR_LETTERHEAD_PATH}")

    raw = TEMPLATE_PATH.read_bytes()
    doc = Document(io.BytesIO(raw))

    c = data.customer
    tables = doc.tables

    if len(tables) >= 1:
        t0: Table = tables[0]

        _append_to_label(t0.rows[0].cells[0], "Date:", _fmt_date(data.request_date))
        _append_to_label(t0.rows[0].cells[1], "Lab Code:", data.lab_code or "")

        _set_cell_text(
            t0.rows[1].cells[0],
            f"Customer Details:\n{c.customer_name or ''}",
        )
        _set_cell_text(
            t0.rows[1].cells[1],
            f"Address:\n{c.address or ''}",
        )

        contact_names, contact_emails = format_contacts_for_display(
            c.resolved_contacts()
        )
        _append_to_label(
            t0.rows[2].cells[0],
            "Name of Contact Person:",
            contact_names or c.contact_person or "",
        )
        _append_to_label(
            t0.rows[2].cells[1],
            "Contact Number:",
            c.contact_number or "",
        )

        _append_to_label(
            t0.rows[3].cells[0],
            "Email:",
            contact_emails or c.email or "",
        )
        _append_to_label(
            t0.rows[3].cells[1],
            "GST Number of Customer:",
            c.gst_number or "",
        )

        nos = "" if data.number_of_samples is None else str(data.number_of_samples)
        _set_cell_text(t0.rows[4].cells[0], "Number of Samples")
        _set_cell_text(t0.rows[4].cells[1], nos)

        filled_preview = [s for s in data.samples if not s.is_empty()]
        first_details = (
            effective_sample_request_details(filled_preview[0], data)
            if filled_preview
            else {}
        )

        _set_cell_text(t0.rows[5].cells[0], "Sampling Done by Laboratory:")
        _set_cell_text(
            t0.rows[5].cells[1],
            _yes_no_cell(first_details.get("sampling_by_lab")),
        )

        _set_cell_text(
            t0.rows[6].cells[0],
            "Storage Temperature of sample required:",
        )
        _set_cell_text(
            t0.rows[6].cells[1],
            str(first_details.get("storage_temperature") or ""),
        )

        method = str(first_details.get("test_method_spec") or "")
        _set_cell_text(
            t0.rows[7].cells[0],
            f"Specific test method/ Specification to be followed:\n{method}".rstrip(),
        )

        _set_cell_text(t0.rows[8].cells[0], "Decision Rule required:")
        _set_cell_text(
            t0.rows[8].cells[1],
            _yes_no_cell(first_details.get("decision_rule")),
        )

        _set_cell_text(t0.rows[9].cells[0], "Service required:")
        _set_cell_text(
            t0.rows[9].cells[1],
            _service_cell(str(first_details.get("service_type") or "")),
        )

        _set_cell_text(t0.rows[10].cells[0], "Mode of report delivery:")
        _set_cell_text(
            t0.rows[10].cells[1],
            _delivery_cell(str(first_details.get("delivery_mode") or "")),
        )

        _set_cell_text(t0.rows[11].cells[0], "Payment Details:")
        _set_cell_text(t0.rows[11].cells[1], data.payment_details or "")

    filled = [s for s in data.samples if not s.is_empty()]

    sample_col_widths = _CTR_DEFAULT_SAMPLE_COL_WIDTHS
    if len(tables) >= 3:
        grid = read_table_grid_column_widths(tables[2])
        if len(grid) == 5:
            sample_col_widths = grid
        remove_table(tables[2])

    remove_body_paragraph_containing(doc, _STALE_SAMPLE_DESCRIPTION_NEEDLE)

    _apply_ctr_signature_block(
        doc,
        generated_by=generated_by,
        generated_at=generated_at,
        request_date=data.request_date,
    )

    for index, sample in enumerate(filled, start=1):
        details = effective_sample_request_details(sample, data)
        storage_temp = str(details.get("storage_temperature") or "")
        if index > 1:
            _add_page_break(doc)
        _add_ctr_bold_heading(doc, CHECKLIST_TITLE)
        _add_verification_table(doc, sample, storage_temperature=storage_temp)
        _add_ctr_bold_heading(doc, SAMPLE_SECTION_HEADING)
        _add_sample_table(
            doc, sample, index, col_widths=sample_col_widths
        )
        _add_tests_paragraphs(doc, sample)

    apply_ctr_header_only(doc)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def suggest_docx_filename(data: TestRequestData) -> str:
    """Filename for the filled Word download."""
    name = (data.customer.customer_name or "customer").strip()
    safe = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in name)
    safe = safe[:40].strip("_") or "customer"
    d = data.request_date.isoformat() if data.request_date else "undated"
    return f"CTR_{safe}_{d}.docx"
