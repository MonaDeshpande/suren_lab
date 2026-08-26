"""
One-off builder: reference/Test Report Format.docx

Clones layout from reference/SLS-26-562-05 DK Brothers Test Report.docx
(A4, Word header band, metadata grid, vMerge results header, Book Antiqua disclaimer).

Run: python scripts/build_food_test_report_template.py
"""

from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

SOURCE = ROOT / "reference" / "SLS-26-562-05 DK Brothers Test Report.docx"
OUT = ROOT / "reference" / "Test Report Format.docx"

# First body element index where page-2 content begins (after page-1 End of Report).
PAGE2_START_INDEX = 22

DISCLAIMER_BULLETS = [
    "Sample submitted by the customer in their own container.",
    "Above analysis result is valid only for specific sample as stated above, "
    "without any bias to its source.",
    "We claim no responsibility for changes made in the report after dispatch "
    "e.g. Use of whitener or eraser.",
    "Above test report cannot be produced as legal evidence without our prior "
    "written permission.",
    "Sample stored for one week and test Report for one year from the date received.",
    "Duplicate copies of Report or Invoice will be charged extra.",
]

REMARK_PLACEHOLDER = (
    "Remark: The sample given for analysis confirm to the specifications of FSSAI 2006, "
    "Rules and Regulations Latest Amendments Up to 12/08/2025 as per the above tests "
    "processed. The Results are pertaining to the sample sent for analysis only. "
    "The Specifications mentioned are for whole Jaggery."
)

METADATA_LABELS = [
    ("Customer Name & Address", "", "", ""),
    ("Customer Sample ID", "", "Batch No.", ""),
    ("Lab Code ", "", "Date of Sample Receipt", ""),
    ("Sample Name", "", "Sample Drawn By", ""),
    ("Condition Of Sample", "", "Test Performance Date", ""),
    ("Tests processed", "", "Sample quantity", ""),
    ("Appearance", "", "", ""),
    ("Date of Sampling", "--", "Sampling Done by", ""),
    ("Location of sampling", "--", "Sampling Method", "--"),
]

RESULTS_HEADERS = [
    "Sr. No",
    "Name of Test",
    "Result",
    "Specifications\nFSSAI 2006, Rules & Regulations, Latest Amendments Up to 12/08/2025",
    "Method of Analysis",
]


def _set_para_text(para: Paragraph, text: str) -> None:
    text = text or ""
    if not para.runs:
        para.text = text
        return
    para.runs[0].text = text
    for run in para.runs[1:]:
        run.text = ""


def _set_cell_text(cell, text: str) -> None:
    text = text or ""
    if not cell.paragraphs:
        cell.text = text
        return
    _set_para_text(cell.paragraphs[0], text)
    for para in cell.paragraphs[1:]:
        _set_para_text(para, "")


def _remove_body_elements_from(doc: Document, start_index: int) -> None:
    body = doc.element.body
    children = list(body)
    for child in children[start_index:]:
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        body.remove(child)


def _clear_metadata_table(table: Table) -> None:
    for i, (l1, v1, l2, v2) in enumerate(METADATA_LABELS):
        if i >= len(table.rows):
            break
        row = table.rows[i]
        _set_cell_text(row.cells[0], l1)
        _set_cell_text(row.cells[1], v1)
        if l2:
            _set_cell_text(row.cells[2], l2)
            _set_cell_text(row.cells[3], v2)
        else:
            _set_cell_text(row.cells[2], "")
            _set_cell_text(row.cells[3], "")


def _clear_results_table(table: Table) -> None:
    if len(table.rows) < 3:
        return
    for col, header in enumerate(RESULTS_HEADERS):
        if col < len(table.rows[0].cells):
            _set_cell_text(table.rows[0].cells[col], header)
        if col < len(table.rows[1].cells):
            _set_cell_text(table.rows[1].cells[col], header)
    for ri in range(2, len(table.rows) - 1):
        for ci in range(min(5, len(table.rows[ri].cells))):
            _set_cell_text(table.rows[ri].cells[ci], "")
    remark_row = table.rows[-1]
    if len(remark_row.cells) > 1:
        _set_cell_text(remark_row.cells[1], REMARK_PLACEHOLDER)


def _reset_body_paragraphs(doc: Document) -> None:
    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if i == 0 or text.startswith("ULR No"):
            _set_para_text(para, "ULR No:")
        elif text.startswith("Date:") or "Report No:" in text:
            _set_para_text(
                para,
                "Date: --/--/----"
                "\t\t\t\t                                  "
                "Report No: ---/26/---/--/--",
            )
        elif text == "CHEMICAL TEST REPORT":
            _set_para_text(para, "CHEMICAL TEST REPORT")
        elif text.startswith("Dr.") or text.startswith("Mrs."):
            _set_para_text(
                para,
                "{{AUTHORIZED_SIGNATORY}}"
                "\t\t\t\t                                  "
                "Checked by: {{CHECKED_BY}}",
            )
        elif "{{AUTHORIZED_SIGNATORY}}" in text:
            _set_para_text(
                para,
                "{{AUTHORIZED_SIGNATORY}}"
                "\t\t\t\t                                  "
                "Checked by: {{CHECKED_BY}}",
            )
        elif text == "Director":
            _set_para_text(para, "Director")
        elif text == "Authorized signatory":
            _set_para_text(para, "Authorized signatory")
        elif text.startswith("For,"):
            _set_para_text(para, "For, SLS")
        elif text.startswith("Checked by"):
            _set_para_text(para, "Checked by: {{CHECKED_BY}}")
        elif text == "Disclaimer" or text.startswith("Disclaimer"):
            _set_para_text(para, "Disclaimer")
        elif text == "End of Report":
            _set_para_text(para, "End of Report")
        elif text in DISCLAIMER_BULLETS or any(
            text.startswith(b[:20]) for b in DISCLAIMER_BULLETS
        ):
            continue
        elif not text:
            continue
        else:
            _set_para_text(para, "")

    disc_started = False
    bullet_idx = 0
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text == "Disclaimer" or text.startswith("Disclaimer"):
            disc_started = True
            _set_para_text(para, "Disclaimer")
            continue
        if disc_started and bullet_idx < len(DISCLAIMER_BULLETS):
            if text == "End of Report":
                break
            _set_para_text(para, DISCLAIMER_BULLETS[bullet_idx])
            bullet_idx += 1


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Reference missing: {SOURCE}")

    doc = Document(str(SOURCE))
    _remove_body_elements_from(doc, PAGE2_START_INDEX)

    if len(doc.tables) >= 1:
        _clear_metadata_table(doc.tables[0])
    if len(doc.tables) >= 2:
        _clear_results_table(doc.tables[1])

    _reset_body_paragraphs(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
