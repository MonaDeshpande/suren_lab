"""
Shared python-docx helpers for validation Word documents.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from validation.inventory import ORGANIZATION, SYSTEM_NAME, SYSTEM_VERSION


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    return doc


def add_cover(
    doc: Document,
    *,
    doc_title: str,
    doc_code: str,
    generated_at: str,
    subtitle: str = "DRAFT – for review and sign-off",
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(ORGANIZATION)
    run.bold = True
    run.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SYSTEM_NAME)
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_title)
    run.bold = True
    run.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.italic = True
    run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    run.font.size = Pt(12)

    doc.add_paragraph()
    meta = [
        ("Document code", doc_code),
        ("System version", SYSTEM_VERSION),
        ("Generated (Asia/Kolkata)", generated_at),
        ("Status", "DRAFT"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(meta):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
        _bold_cell(table.rows[i].cells[0])

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        "This document was auto-generated from the S_LAB codebase inventory. "
        "It is a draft for SME review. Do not treat blank Pass/Fail fields as executed results."
    ).italic = True
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    col_widths: Optional[Sequence[float]] = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        hdr.cells[i].text = h
        _bold_cell(hdr.cells[i])
        _shade_cell(hdr.cells[i], "D9E2F3")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = "" if val is None else str(val)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
    doc.add_paragraph()


def add_signature_block(doc: Document) -> None:
    add_heading(doc, "Document approval", level=1)
    add_para(
        doc,
        "Complete after review. Electronic or wet-ink signatures are acceptable per lab QMS.",
    )
    headers = ("Role", "Name", "Signature", "Date")
    rows = [
        ("Author / Generator operator", "", "", ""),
        ("System owner / Lab manager", "", "", ""),
        ("QA / Quality reviewer", "", "", ""),
        ("IT / Validation (if applicable)", "", "", ""),
    ]
    add_table(doc, headers, rows)


def save_document(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def dated_filename(prefix: str, stamp: Optional[str] = None) -> str:
    if not stamp:
        stamp = datetime.now().strftime("%Y%m%d")
    return f"{prefix}_{stamp}.docx"


def _bold_cell(cell) -> None:
    if not cell.paragraphs:
        return
    p = cell.paragraphs[0]
    text = p.text
    p.clear()
    run = p.add_run(text)
    run.bold = True


def _shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)
