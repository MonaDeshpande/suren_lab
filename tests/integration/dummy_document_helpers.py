"""
Shared helpers for integration tests that generate CTR/protocol files and wipe DB rows.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from db.connection import get_db
from services.customers import Customer

INTEGRATION_PREVIEW_ROOT = (
    Path(__file__).resolve().parent.parent / "output_preview" / "integration"
)


def ensure_integration_preview_dirs() -> dict[str, Path]:
    """Create integration preview subdirs; return mapping category -> path."""
    subdirs = ("ctr", "water", "food", "micro")
    out: dict[str, Path] = {}
    for name in subdirs:
        path = INTEGRATION_PREVIEW_ROOT / name
        path.mkdir(parents=True, exist_ok=True)
        out[name] = path
    return out


def dummy_customer(gst: str, *, name: str = "Pytest Dummy Customer Pvt Ltd") -> Customer:
    return Customer(
        customer_name=name,
        address="1 Integration Test Road, Nashik",
        contact_person="Test User",
        contact_number="9876500000",
        email="pytest-dummy@example.com",
        gst_number=gst,
    )


def write_preview_bytes(out_dir: Path, filename: str, data: bytes) -> Path:
    safe = filename.replace("/", "-").replace("\\", "-")
    path = out_dir / safe
    path.write_bytes(data)
    return path


def write_protocol_preview(
    out_dir: Path,
    *,
    docx_filename: str,
    docx_bytes: bytes,
    pdf_bytes: bytes | None,
) -> tuple[Path, Path | None]:
    docx_path = write_preview_bytes(out_dir, docx_filename, docx_bytes)
    pdf_path = None
    if pdf_bytes:
        pdf_path = out_dir / (Path(docx_filename.replace("/", "-")).stem + ".pdf")
        pdf_path.write_bytes(pdf_bytes)
    return docx_path, pdf_path


def pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def docx_body_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def assert_db_clean_for_gst(gst: str) -> None:
    normalized = (gst or "").strip().upper()
    with get_db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT 1 FROM customers WHERE gst_number = %s LIMIT 1",
                (normalized,),
            )
            assert cur.fetchone() is None
