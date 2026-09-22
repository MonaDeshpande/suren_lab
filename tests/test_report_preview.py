"""
Headless CTR preview: DB round-trip → PDF + DOCX written to tests/output_preview/ctr/.

Run: pytest tests/test_report_preview.py -v -s
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from services.ctr_pdf import generate_ctr_documents
from services.customers import Customer
from services.docx_filler import suggest_docx_filename
from services.pdf_generator import ctr_total_pages, suggest_pdf_filename
from services.requests import SampleRow, TestRequestData, get_test_request, save_test_request
from tests.conftest import cleanup_test_data, sample_verification_kwargs
from tests.test_db_integration import build_multi_type_ctr_data

OUTPUT_DIR = Path(__file__).resolve().parent / "output_preview" / "ctr"

CTR_TABLE_HEADERS = [
    "Sr. No",
    "Name of sample",
    "Code/batch no.",
    "Sample qty.",
    "Parameters",
]


@pytest.fixture(scope="session")
def preview_output_dir() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR


def _write_preview(out_dir: Path, filename: str, data: bytes) -> Path:
    path = out_dir / filename.replace("/", "-").replace("\\", "-")
    path.write_bytes(data)
    return path


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_body(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _assert_ctr_content(pdf_bytes: bytes, docx_bytes: bytes, *, customer_name: str) -> None:
    assert len(pdf_bytes) > 5000
    assert len(docx_bytes) > 5000

    pdf_text = _pdf_text(pdf_bytes)
    assert customer_name in pdf_text
    assert "Sample Verification Checklist" in pdf_text
    assert "✓" in pdf_text
    assert "Yes ( X )" not in pdf_text

    docx_text = _docx_body(docx_bytes)
    assert customer_name in docx_text
    assert "✓" in docx_text
    header_found = any(
        [c.text.strip() for c in table.rows[0].cells] == CTR_TABLE_HEADERS
        for table in Document(io.BytesIO(docx_bytes)).tables
    )
    assert header_found


@pytest.mark.integration
def test_ctr_multi_sample_document_bundle_writes_preview(
    require_db,
    qa_analyst_pair,
    db_test_gst,
    preview_output_dir,
):
    """TC-REC-037: multi-sample CTR PDF+DOCX via generate_ctr_documents → output_preview/ctr."""
    chem, micro = qa_analyst_pair
    data = build_multi_type_ctr_data(chem.id, micro.id, gst=db_test_gst)
    saved = save_test_request(data, actor=None)
    try:
        loaded = get_test_request(saved.request_id)
        assert loaded is not None
        filled = [s for s in loaded.samples if not s.is_empty()]
        expected_pages = ctr_total_pages(loaded)

        docx_bytes, pdf_bytes, docx_name, pdf_name, pdf_err = generate_ctr_documents(
            loaded,
            generated_by="CTR Preview Test",
            generated_at="27/07/2026 10:00",
        )
        assert pdf_err is None
        assert pdf_bytes is not None
        _assert_ctr_content(
            pdf_bytes,
            docx_bytes,
            customer_name="ABC Foods Pvt Ltd",
        )

        reader = PdfReader(io.BytesIO(pdf_bytes))
        assert len(reader.pages) == expected_pages
        pdf_text = _pdf_text(pdf_bytes)
        for name in (
            "Potable Water",
            "Paneer Gravy",
            "Jaggery",
            "Cattle Feed Mix",
        ):
            assert name in pdf_text
        assert "Prepared by:" in pdf_text
        assert f"Page 1 of {expected_pages}" in pdf_text
        assert f"Page {expected_pages} of {expected_pages}" in pdf_text

        pdf_path = _write_preview(preview_output_dir, pdf_name, pdf_bytes)
        docx_path = _write_preview(preview_output_dir, docx_name, docx_bytes)
        print(f"\nCTR multi-sample preview written to:\n  {pdf_path}\n  {docx_path}")
        assert len(filled) == 5
    finally:
        cleanup_test_data(gst=db_test_gst)


@pytest.mark.integration
def test_ctr_preview_from_db_round_trip(
    require_db,
    qa_analyst_pair,
    db_test_gst,
    preview_output_dir,
):
    """TC-REC-037 automation: save → reload → generate CTR artifacts for review."""
    chem, micro = qa_analyst_pair
    data = build_multi_type_ctr_data(chem.id, micro.id, gst=db_test_gst)
    saved = save_test_request(data, actor=None)
    try:
        loaded = get_test_request(saved.request_id)
        assert loaded is not None

        docx_bytes, pdf_bytes, _, _, pdf_err = generate_ctr_documents(loaded)
        assert pdf_err is None
        assert pdf_bytes is not None
        _assert_ctr_content(
            pdf_bytes,
            docx_bytes,
            customer_name="ABC Foods Pvt Ltd",
        )

        for name in (
            "Potable Water",
            "Paneer Gravy",
            "Jaggery",
        ):
            assert name in _pdf_text(pdf_bytes)

        pdf_path = _write_preview(
            preview_output_dir,
            suggest_pdf_filename(loaded),
            pdf_bytes,
        )
        docx_path = _write_preview(
            preview_output_dir,
            suggest_docx_filename(loaded),
            docx_bytes,
        )

        print(f"\nCTR preview written to:\n  {pdf_path}\n  {docx_path}")
    finally:
        cleanup_test_data(gst=db_test_gst)


def test_ctr_preview_food_multi_sample_from_memory(preview_output_dir):
    """Unit-speed preview (no DB): sauce + honey rows from docx_filler tests."""
    data = TestRequestData(
        customer=Customer(
            customer_name="ABC Foods",
            address="1 Lab Road",
            contact_person="Ravi",
            contact_number="9876543210",
            email="ravi@example.com",
            gst_number="27AAAAA0000A1Z5",
        ),
        request_date=date(2026, 7, 30),
        lab_code="LAB/CTR/26/001",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="sauce",
                batch_code="01",
                quantity="100gm",
                parameters="FSSAI",
                test_keys=["moisture"],
                package_type="fssai",
                **sample_verification_kwargs(verify_lab_code="LAB/CTR/26/001"),
            ),
            SampleRow(
                sr_no=2,
                sample_name="honey",
                batch_code="H-02",
                quantity="250 g",
                parameters="Basic Nutrition",
                test_keys=["bn_protein", "bn_total_fat"],
                package_type="nutrition_only",
                **sample_verification_kwargs(verify_lab_code="LAB/CTR/26/001"),
            ),
        ],
    )

    docx_bytes, pdf_bytes, _, _, pdf_err = generate_ctr_documents(data)
    assert pdf_err is None
    assert pdf_bytes is not None
    _assert_ctr_content(pdf_bytes, docx_bytes, customer_name="ABC Foods")
    assert "sauce" in _pdf_text(pdf_bytes)
    assert "honey" in _pdf_text(pdf_bytes)

    pdf_path = _write_preview(
        preview_output_dir,
        suggest_pdf_filename(data),
        pdf_bytes,
    )
    docx_path = _write_preview(
        preview_output_dir,
        suggest_docx_filename(data),
        docx_bytes,
    )
    print(f"\nCTR food preview written to:\n  {pdf_path}\n  {docx_path}")
