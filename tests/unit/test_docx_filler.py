"""Unit tests for CTR DOCX filler — per-sample pages and verification checklist."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from services.customers import Customer
from services.document_templates import CTR_FORM_BODY_PATH
from services.docx_filler import fill_docx_bytes
from services.ctr_verification import CHECKLIST_TITLE, SAMPLE_SECTION_HEADING
from services.requests import SampleRow, TestRequestData
from tests.conftest import sample_verification_kwargs


def _ctr_request() -> TestRequestData:
    return TestRequestData(
        customer=Customer(
            customer_name="ABC Foods",
            address="1 Lab Road",
            contact_person="Ravi",
            contact_number="9876543210",
            email="ravi@example.com",
            gst_number="27AAAAA0000A1Z5",
        ),
        request_date=date(2026, 7, 30),
        lab_code="LAB/CTR/26/001",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Jaggery",
                batch_code="B-01",
                quantity="500 g",
                parameters="FSSAI",
                test_keys=["moisture", "total_ash"],
                package_type="fssai",
                **sample_verification_kwargs(verify_lab_code="LAB/CTR/26/001"),
            ),
            SampleRow(
                sr_no=2,
                sample_name="Honey",
                batch_code="H-02",
                quantity="250 g",
                parameters="Basic Nutrition",
                test_keys=["bn_protein", "bn_total_fat"],
                package_type="nutrition_only",
                **sample_verification_kwargs(verify_lab_code="LAB/CTR/26/001"),
            ),
        ],
    )


def _header_text(doc: Document) -> str:
    header = doc.sections[0].header
    parts: list[str] = []
    for table in header.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text or "")
    for para in header.paragraphs:
        parts.append(para.text or "")
    return "\n".join(parts)


def _footer_approval_cells(doc: Document) -> tuple[str, str, str]:
    footer = doc.sections[0].footer
    if not footer.tables:
        return "", "", ""
    table = footer.tables[0]
    prepared = table.rows[0].cells[1].text.strip() if len(table.rows) >= 1 else ""
    reviewed = table.rows[1].cells[1].text.strip() if len(table.rows) >= 2 else ""
    reviewed_label = table.rows[1].cells[0].text.strip() if len(table.rows) >= 2 else ""
    return prepared, reviewed, reviewed_label


def _footer_instr_text(doc: Document) -> list[str]:
    footer_el = doc.sections[0].footer._element
    return [
        (node.text or "").strip()
        for node in footer_el.iter(qn("w:instrText"))
        if (node.text or "").strip()
    ]


class TestDocxSampleTable:
    def test_page_one_has_form_tables_only(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        assert len(doc.tables[0].columns) == 2
        assert len(doc.tables[1].columns) == 1
        for table in doc.tables[:2]:
            assert len(table.columns) != 5

    def test_storage_temperature_in_right_column(self):
        data = _ctr_request()
        data.samples[0].storage_temperature = "25°C"
        doc = Document(BytesIO(fill_docx_bytes(data)))
        t0 = doc.tables[0]
        assert "Storage Temperature" in t0.rows[6].cells[0].text
        assert t0.rows[6].cells[1].text.strip() == "25°C"
        assert "25°C" not in t0.rows[6].cells[0].text

    def test_number_of_samples_in_right_column(self):
        data = _ctr_request()
        data.number_of_samples = 3
        doc = Document(BytesIO(fill_docx_bytes(data)))
        t0 = doc.tables[0]
        assert t0.rows[4].cells[0].text.strip() == "Number of Samples"
        assert t0.rows[4].cells[1].text.strip() == "3"
        assert "3" not in t0.rows[4].cells[0].text

    def test_signature_block_shows_receiver_and_reception_name(self):
        doc = Document(
            BytesIO(
                fill_docx_bytes(
                    _ctr_request(),
                    generated_by="Priya Reception",
                    generated_at="2026-09-22 14:30",
                )
            )
        )
        body = "\n".join(p.text for p in doc.paragraphs)
        assert "Receiver's Sign" not in body
        sign_tables = [
            t
            for t in doc.tables
            if len(t.rows) == 3
            and len(t.columns) == 2
            and t.rows[2].cells[0].text.strip() == "Receiver"
        ]
        assert len(sign_tables) == 1
        assert sign_tables[0].rows[0].cells[0].text.strip() == "Priya Reception"
        assert sign_tables[0].rows[0].cells[1].text.strip().startswith(
            "Customer Signature"
        )
        assert sign_tables[0].rows[1].cells[0].text.strip() == "22/09/2026"

    def test_page_break_before_second_sample_only(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        page_breaks = 0
        for para in doc.paragraphs:
            for run in para.runs:
                for el in run._element:
                    if el.tag == qn("w:br") and el.get(qn("w:type")) == "page":
                        page_breaks += 1
        assert page_breaks == 1

    def test_tests_to_be_performed_heading_is_bold(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        matches = [
            p
            for p in doc.paragraphs
            if (p.text or "").strip() == "Tests to be performed:"
            and p.runs
            and p.runs[0].bold
        ]
        assert len(matches) == 2

    def test_signature_table_before_checklist_tables(self):
        doc = Document(
            BytesIO(
                fill_docx_bytes(_ctr_request(), generated_by="Reception User")
            )
        )
        sign_idx = None
        checklist_idx = None
        for i, table in enumerate(doc.tables):
            if (
                len(table.rows) == 3
                and len(table.columns) == 2
                and table.rows[2].cells[0].text.strip() == "Receiver"
            ):
                sign_idx = i
            if (
                len(table.rows) >= 12
                and table.rows[0].cells[1].text.strip() == "Particulars"
            ):
                if checklist_idx is None:
                    checklist_idx = i
        assert sign_idx is not None
        assert checklist_idx is not None
        assert sign_idx < checklist_idx

    def test_sample_section_headings_are_bold(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        bold_titles = [
            p
            for p in doc.paragraphs
            if (p.text or "").strip() in (CHECKLIST_TITLE, SAMPLE_SECTION_HEADING)
            and p.runs
            and p.runs[0].bold
        ]
        assert CHECKLIST_TITLE in {p.text.strip() for p in bold_titles}
        assert SAMPLE_SECTION_HEADING in {p.text.strip() for p in bold_titles}

    def test_sample_five_column_table_not_on_page_one(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        page_one_five_col = [
            t
            for t in doc.tables[:2]
            if len(t.columns) == 5
        ]
        assert page_one_five_col == []
        sample_tables = [
            t
            for t in doc.tables
            if len(t.rows) >= 2
            and len(t.columns) == 5
            and t.rows[0].cells[0].text.strip() == "Sr. No"
        ]
        assert len(sample_tables) == 2
        assert sample_tables[0].rows[1].cells[1].text.strip() == "Jaggery"

    def test_stale_sample_description_paragraph_removed(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        stale_lines = [
            (p.text or "").strip()
            for p in doc.paragraphs
            if "tests to be perform" in (p.text or "").lower()
            and not (p.text or "").strip().endswith("performed:")
        ]
        assert stale_lines == []

    def test_second_sample_on_later_table_with_tests(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        body_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Tests to be performed:" in body_text
        assert "1. Moisture" in body_text
        assert "2. Total ash on dry basis" in body_text
        assert "1. Protein" in body_text
        honey_tables = [
            t
            for t in doc.tables
            if len(t.rows) >= 2 and t.rows[1].cells[1].text.strip() == "Honey"
        ]
        assert len(honey_tables) == 1

    def test_verification_checklist_tables_appended(self):
        doc = Document(BytesIO(fill_docx_bytes(_ctr_request())))
        checklist_tables = [
            t
            for t in doc.tables
            if len(t.rows) >= 12
            and t.rows[0].cells[1].text.strip() == "Particulars"
        ]
        assert len(checklist_tables) == 2
        assert checklist_tables[0].rows[4].cells[1].text.strip() == (
            "Storage Temperature"
        )
        assert checklist_tables[0].rows[6].cells[1].text.strip() == (
            "Checked for Sample Quantity"
        )
        assert "Yes ( ✓ )" in checklist_tables[0].rows[6].cells[2].text

    def test_no_generator_stamp_in_ctr_docx(self):
        doc = Document(
            BytesIO(
                fill_docx_bytes(
                    _ctr_request(),
                    generated_by="Reception Test",
                    generated_at="2026-07-31 12:24",
                )
            )
        )
        body_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Generated by:" not in body_text
        assert "Date / time:" not in body_text

        header_text = _header_text(doc)
        assert "CUSTOMER TEST REQUEST FORM" in header_text

        prepared, reviewed_value, reviewed_label = _footer_approval_cells(doc)
        assert prepared == ""
        assert reviewed_value == ""
        assert reviewed_label == ""

        instr = _footer_instr_text(doc)
        assert not any("PAGE" in x for x in instr)

        header_el = doc.sections[0].header._element
        header_instr = [
            (node.text or "").strip()
            for node in header_el.iter(qn("w:instrText"))
            if (node.text or "").strip()
        ]
        assert any("NUMPAGES" in x for x in header_instr)

        llp = Document(BytesIO(CTR_FORM_BODY_PATH.read_bytes()))
        assert doc.sections[0].footer_distance == llp.sections[0].footer_distance
