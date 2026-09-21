"""Unit tests for Micro analyst protocol DOCX filler."""

from __future__ import annotations

import json
from datetime import date
from io import BytesIO

import pytest
from docx import Document

from services.micro_protocol_docx import (
    MICRO_PROTOCOL_TEMPLATE,
    fill_micro_protocol_docx_bytes,
    suggest_micro_protocol_filename,
)
from services.protocol_store import ProtocolHeader, TestResultRow
from services.protocols.test_catalog import MICRO_TEST_KEYS
from services.samples import SampleRecord


def _sample(**overrides) -> SampleRecord:
    base = dict(
        id=1,
        request_id=1,
        sample_code="SLS-260704-0001",
        sr_no=1,
        sample_name="Cooked Food: Paneer Gravy",
        batch_code="",
        quantity="250 gm",
        parameters="Paneer Gravy",
        tests_to_perform="",
        status="in_progress",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        category="micro",
        tests_json=json.dumps(list(MICRO_TEST_KEYS)),
        lab_code="SLS/26/546/03",
        customer_name="Test Customer",
        customer_address="Pune",
        report_format="with_logo",
        tests_with_logo_json="",
        tests_without_logo_json="",
        sampling_by_lab=False,
    )
    base.update(overrides)
    return SampleRecord(**base)


def _header() -> ProtocolHeader:
    return ProtocolHeader(
        sample_id=1,
        protocol_no="P-MICRO-001",
        issued_to="Analyst User",
        issued_by="Reception",
        sample_received_on=date(2026, 7, 4),
        date_of_analysis=date(2026, 7, 9),
        appearance_text="",
    )


@pytest.mark.skipif(
    not MICRO_PROTOCOL_TEMPLATE.exists(),
    reason="reference/Micro Protocol.docx missing",
)
class TestMicroProtocolDocx:
    def test_suggest_filename(self):
        assert suggest_micro_protocol_filename(_sample()) == (
            "Micro_Protocol_SLS-260704-0001.docx"
        )

    def test_fill_header_and_results_table(self):
        results = [
            TestResultRow(
                test_key="total_plate_count",
                test_name="Total Plate Count",
                method="IS:5402:2018",
                unit="cfu/gm",
                inputs={"result_value": "3.0 x 10³", "result_unit": "cfu/gm"},
                result_value="3.0 x 10³",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="e_coli",
                test_name="E. Coli",
                method="IS 5887 (Part - 1 ) : 1976 RA 2018",
                unit="cfu/25g",
                inputs={"result_value": "Absent", "result_unit": "cfu/25g"},
                result_value="Absent",
                result_numeric=None,
            ),
        ]
        docx_bytes = fill_micro_protocol_docx_bytes(
            _sample(), _header(), results, generated_by="Analyst"
        )
        assert docx_bytes[:2] == b"PK"
        doc = Document(BytesIO(docx_bytes))
        assert len(doc.tables) >= 3

        t0 = doc.tables[0]
        assert "P-MICRO-001" in t0.rows[0].cells[1].text
        assert "Analyst User" in t0.rows[0].cells[3].text

        t1 = doc.tables[1]
        assert "Paneer Gravy" in t1.rows[0].cells[1].text
        assert "04/07/2026" in t1.rows[0].cells[3].text
        assert "SLS-260704-0001" in t1.rows[1].cells[1].text
        assert "09/07/2026" in t1.rows[1].cells[3].text

        t2 = doc.tables[2]
        assert t2.rows[2].cells[0].text.strip() == "1"
        assert "Total Plate Count" in t2.rows[2].cells[1].text
        result_cell = t2.rows[2].cells[2].text
        assert "3.0 x 10" in result_cell and "cfu/gm" in result_cell
        assert "Absent" not in result_cell
        assert "IS:5402:2018" in t2.rows[2].cells[3].text
        assert "Absent cfu/25g" in t2.rows[4].cells[2].text
