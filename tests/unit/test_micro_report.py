"""Unit tests for Micro category catalog and final report DOCX."""

from __future__ import annotations

import json
from datetime import date

from docx import Document

from services.micro_report_catalog import (
    MICRO_REPORT_SPECS,
    micro_report_keys_ordered,
    spec_for_key,
)
from services.protocol_store import ProtocolHeader, TestResultRow
from services.protocols.test_catalog import (
    CATEGORY_MICRO,
    MICRO_TEST_KEYS,
    SAMPLE_CATEGORIES,
    TEST_CATALOG,
    catalog_keys_for_category,
    default_test_keys_for_category,
    get_test,
    normalize_category,
)
from services.samples import SampleRecord
from services.test_report_micro_docx import (
    MicroReportFillOptions,
    fill_micro_test_report_docx_bytes,
)
from services.test_report_pdf import generate_final_report


def test_micro_category_registered():
    assert CATEGORY_MICRO in SAMPLE_CATEGORIES
    assert SAMPLE_CATEGORIES[CATEGORY_MICRO] == "Micro"
    assert normalize_category("micro") == CATEGORY_MICRO


def test_micro_fixed_panel():
    assert len(MICRO_TEST_KEYS) == 6
    assert catalog_keys_for_category(CATEGORY_MICRO) == MICRO_TEST_KEYS
    assert default_test_keys_for_category(CATEGORY_MICRO) == MICRO_TEST_KEYS
    for key in MICRO_TEST_KEYS:
        test = get_test(key)
        assert CATEGORY_MICRO in test.categories
        assert test.key in TEST_CATALOG
        assert key in MICRO_REPORT_SPECS
        assert spec_for_key(key).limits
        assert spec_for_key(key).method


def test_micro_result_calculator():
    test = get_test("e_coli")
    display, numeric = test.calculate(
        {"result_value": "Absent", "result_unit": "cfu/25g"},
        {},
    )
    assert display == "Absent"
    assert numeric is None


def test_micro_report_keys_ordered():
    assert micro_report_keys_ordered(None) == MICRO_TEST_KEYS
    assert micro_report_keys_ordered({"e_coli", "salmonella"}) == [
        "e_coli",
        "salmonella",
    ]


def _sample(**overrides) -> SampleRecord:
    base = dict(
        id=1,
        request_id=1,
        sample_code="SLS-260704-0001",
        sr_no=1,
        sample_name="Cooked Food: Paneer Gravy",
        batch_code="",
        quantity="250 gm",
        parameters="Paneer Gravy",
        tests_to_perform="",
        status="completed",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        category="micro",
        tests_json=json.dumps(list(MICRO_TEST_KEYS)),
        lab_code="SLS/26/546/03",
        customer_name="Test Customer",
        customer_address="Pune",
        report_format="with_logo",
        tests_with_logo_json="",
        tests_without_logo_json="",
        sampling_by_lab=False,
    )
    base.update(overrides)
    return SampleRecord(**base)


def _header() -> ProtocolHeader:
    return ProtocolHeader(
        sample_id=1,
        protocol_no="P-MICRO-001",
        issued_to="Analyst",
        issued_by="Reception",
        sample_received_on=date(2026, 7, 4),
        date_of_analysis=date(2026, 7, 9),
        appearance_text="Reddish orange coloured Paneer gravy.",
    )


def test_fill_micro_report_docx():
    results = [
        TestResultRow(
            test_key="total_plate_count",
            test_name="Total Plate Count",
            method="IS:5402:2018",
            unit="",
            inputs={"result_obs": "3.0 x 10³ cfu/gm"},
            result_value="3.0 x 10³ cfu/gm",
            result_numeric=None,
        ),
        TestResultRow(
            test_key="e_coli",
            test_name="E. Coli",
            method="IS 5887 (Part - 1 ) : 1976 RA 2018",
            unit="",
            inputs={"result_obs": "Absent"},
            result_value="Absent",
            result_numeric=None,
        ),
    ]
    opts = MicroReportFillOptions(
        report_no="SLS/26/546/03/01",
        report_date=date(2026, 7, 9),
        condition_of_sample="Company packing",
        customer_sample_id="Paneer Gravy",
        generated_by="Reviewer",
    )
    docx_bytes = fill_micro_test_report_docx_bytes(
        _sample(), _header(), results, opts=opts
    )
    assert docx_bytes[:2] == b"PK"
    doc = Document(__import__("io").BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text
    assert "Test Report" in full_text or "TEST REPORT" in full_text
    assert "QSF No -7.8.2" in full_text
    assert "SLS/26/546/03/01" in full_text
    assert len(doc.tables) >= 3
    results_table = doc.tables[2]
    assert len(results_table.rows) == 2 + len(MICRO_TEST_KEYS)
    cell_texts = [
        results_table.rows[i].cells[1].text.strip()
        for i in range(2, len(results_table.rows))
    ]
    assert "Total Plate Count" in cell_texts
    assert "E. Coli" in cell_texts
    assert "Shall be Absent" in results_table.rows[4].cells[3].text
    assert "3.0 x 10³ cfu/gm" in results_table.rows[2].cells[2].text or "3.0 x 10" in results_table.rows[2].cells[2].text
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Generated by" not in body_text
    assert "Date / time" not in body_text


def test_micro_default_report_no_is_sample_code():
    sample = _sample(sample_code="SLS/26/546/03")
    docx_bytes = fill_micro_test_report_docx_bytes(sample, _header(), [])
    doc = Document(__import__("io").BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text
    assert "SLS/26/546/03" in full_text
    assert "SLS/26/546/03/01" not in full_text


def test_generate_final_report_routes_micro(monkeypatch):
    """Route Micro → DOCX filler; skip Word→PDF in unit tests (COM hangs on CI)."""
    from services import test_report_micro_docx as micro_mod

    monkeypatch.setattr(
        micro_mod,
        "try_convert_docx_to_pdf",
        lambda _b: None,
    )
    results = [
        TestResultRow(
            test_key="salmonella",
            test_name="Salmonella",
            method=spec_for_key("salmonella").method,
            unit="",
            inputs={"result_obs": "Absent"},
            result_value="Absent",
            result_numeric=None,
        )
    ]
    output = generate_final_report(
        _sample(),
        _header(),
        results,
        generated_by="Reviewer",
        micro_opts=MicroReportFillOptions(condition_of_sample="Sealed"),
    )
    assert output.is_micro is True
    assert output.is_water is False
    assert output.docx_bytes
    assert output.docx_filename.endswith(".docx")
    assert output.pdf_bytes is None
