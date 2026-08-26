"""Unit tests for water final report DOCX generation."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from docx.oxml.ns import qn

from services.protocol_store import ProtocolHeader, TestResultRow
from services.samples import SampleRecord, water_report_keys
from services.test_report_water_docx import (
    WATER_REPORT_TEMPLATE_PATH,
    WaterReportFillOptions,
    fill_water_test_report_docx_bytes,
)
from services.water_report_catalog import (
    WATER_MICRO_PLACEHOLDERS,
    WATER_REPORT_LIMITS,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent


def _sample(**overrides) -> SampleRecord:
    base = dict(
        id=1,
        request_id=1,
        sample_code="SLS-260721-0002",
        sr_no=1,
        sample_name="Potable Water",
        batch_code="---",
        quantity="1 lit",
        parameters="Drinking Water",
        tests_to_perform="",
        status="pending",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        category="water",
        tests_json=json.dumps(
            [
                "ph",
                "tds",
                "chlorides",
                "total_alkalinity",
                "conductivity",
                "total_hardness",
                "calcium_ca",
                "calcium_caco3",
                "magnesium",
                "odor",
                "turbidity",
            ]
        ),
        lab_code="SLS/26/638/01",
        customer_name="Test Customer",
        customer_address="Pune",
        report_format="with_logo",
        tests_with_logo_json="",
        tests_without_logo_json="",
    )
    base.update(overrides)
    return SampleRecord(**base)


def _header() -> ProtocolHeader:
    return ProtocolHeader(
        sample_id=1,
        protocol_no="P-001",
        issued_to="Analyst",
        issued_by="Reception",
        sample_received_on=date(2026, 7, 27),
        date_of_analysis=date(2026, 7, 30),
        appearance_text="Colorless transparent liquid",
    )


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


@pytest.mark.skipif(
    not WATER_REPORT_TEMPLATE_PATH.exists(),
    reason="Water test report template missing",
)
class TestWaterReportDocx:
    def test_calcium_caco3_excluded_from_report_keys(self):
        keys = water_report_keys(_sample())
        assert "calcium_caco3" not in keys
        assert "calcium_ca" in keys

    def test_generates_docx_with_chemical_result(self):
        results = [
            TestResultRow(
                test_key="ph",
                test_name="pH",
                method="IS 3025 : Part 11",
                unit="",
                inputs={},
                result_value="7.320",
                result_numeric=7.32,
            )
        ]
        out = fill_water_test_report_docx_bytes(_sample(), _header(), results)
        doc = Document(io.BytesIO(out))
        chem_table = doc.tables[1]
        assert "07.32" in _cell_text(chem_table.rows[3].cells[2])
        assert WATER_REPORT_LIMITS["ph"].desirable in _cell_text(chem_table.rows[3].cells[3])

    def test_micro_rows_filled_from_protocol(self):
        results = [
            TestResultRow(
                test_key="water_total_coliform",
                test_name="Total Coliform",
                method="",
                unit="",
                inputs={"result_obs": "absent"},
                result_value="absent",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="water_e_coli",
                test_name="E. coli",
                method="",
                unit="",
                inputs={"result_obs": "present"},
                result_value="present",
                result_numeric=None,
            ),
        ]
        out = fill_water_test_report_docx_bytes(_sample(), _header(), results)
        doc = Document(io.BytesIO(out))
        micro_table = doc.tables[4]
        assert _cell_text(micro_table.rows[2].cells[2]) == "Absent"
        assert _cell_text(micro_table.rows[3].cells[2]) == "Present"

    def test_footer_static_one_of_one(self):
        out = fill_water_test_report_docx_bytes(_sample(), _header(), [])
        doc = Document(io.BytesIO(out))
        footer = doc.sections[0].footer
        assert footer.paragraphs[0].text.strip() == "page 1 of 1"
        assert not footer._element.findall(".//" + qn("w:instrText"))

    def test_ulr_hidden_without_logo(self):
        opts = WaterReportFillOptions(ulr_no="TC1611826000063801F")
        out = fill_water_test_report_docx_bytes(
            _sample(), _header(), [], opts=opts, with_logo=False
        )
        doc = Document(io.BytesIO(out))
        assert doc.paragraphs[0].text.strip() == ""

    def test_ulr_shown_with_logo(self):
        opts = WaterReportFillOptions(ulr_no="TC1611826000063801F")
        out = fill_water_test_report_docx_bytes(
            _sample(), _header(), [], opts=opts, with_logo=True
        )
        doc = Document(io.BytesIO(out))
        assert "TC1611826000063801F" in doc.paragraphs[0].text

    def test_two_sections_after_fill(self):
        out = fill_water_test_report_docx_bytes(_sample(), _header(), [])
        doc = Document(io.BytesIO(out))
        assert len(doc.sections) == 2
        assert doc.sections[1].footer.paragraphs[0].text.strip() == "page 1 of 1"

    def test_appearance_and_customer_info_filled(self):
        opts = WaterReportFillOptions(
            condition_of_sample="Packed in plastic container",
            customer_sample_id="Drinking Water",
        )
        out = fill_water_test_report_docx_bytes(
            _sample(), _header(), [], opts=opts
        )
        doc = Document(io.BytesIO(out))
        info = doc.tables[0]
        assert "Colorless transparent liquid" in _cell_text(info.rows[7].cells[1])
        assert "Packed in plastic container" in _cell_text(info.rows[6].cells[1])
        assert "Drinking Water" in _cell_text(info.rows[1].cells[1])

    def test_physical_section_fills_odor(self):
        results = [
            TestResultRow(
                test_key="odor",
                test_name="Odor",
                method="IS 3025 : Part 5",
                unit="",
                inputs={},
                result_value="Agreeable",
                result_numeric=None,
            )
        ]
        out = fill_water_test_report_docx_bytes(_sample(), _header(), results)
        doc = Document(io.BytesIO(out))
        physical = doc.tables[3]
        assert "Agreeable" in _cell_text(physical.rows[3].cells[2])

    def test_chem_and_micro_report_no_match_sample_code(self):
        out = fill_water_test_report_docx_bytes(_sample(), _header(), [])
        doc = Document(io.BytesIO(out))
        joined = "\n".join(p.text for p in doc.paragraphs)
        expected = _sample().sample_code
        assert expected in joined
        assert f"{_sample().lab_code}/02" not in joined
