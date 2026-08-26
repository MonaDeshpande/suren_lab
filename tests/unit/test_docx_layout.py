"""Unit tests for shared DOCX layout helpers."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from services.docx_layout import (
    REPORT_QSF_LABEL,
    finalize_docx_document,
    rewrite_unicode_scripts_in_document,
    set_report_footer,
    set_report_qsf_label,
)
from docx.oxml.ns import qn


def test_rewrite_caco3_subscript_to_vert_align():
    doc = Document()
    doc.add_paragraph("Total Alkalinity as CaCO₃")
    rewrite_unicode_scripts_in_document(doc)
    para = doc.paragraphs[0]
    assert para.text == "Total Alkalinity as CaCO3"
    assert "₃" not in para.text
    sub_runs = [r for r in para.runs if r.font.subscript]
    assert sub_runs
    assert any("3" in (r.text or "") for r in sub_runs)


def test_set_report_qsf_label_normalizes_variants():
    doc = Document()
    doc.add_paragraph("QSP No-7.8.2")
    set_report_qsf_label(doc)
    assert doc.paragraphs[0].text == REPORT_QSF_LABEL


def test_finalize_docx_document_sets_qsf_and_rewrites_scripts():
    doc = Document()
    doc.add_paragraph("QSF No. 7.8.2")
    doc.add_paragraph("Hardness as CaCO₃")
    finalize_docx_document(doc, set_qsf=True)
    assert doc.paragraphs[0].text == REPORT_QSF_LABEL
    assert doc.paragraphs[1].text == "Hardness as CaCO3"
    out = BytesIO()
    doc.save(out)
    assert out.getvalue()[:2] == b"PK"


def _footer_field_count(doc: Document) -> int:
    footer = doc.sections[0].footer
    return len(footer._element.findall(".//" + qn("w:fldChar")))


def test_set_report_footer_static_one_of_one():
    doc = Document()
    section = doc.sections[0]
    footer = section.footer
    footer.add_paragraph("page 1 of 1")
    set_report_footer(doc, with_logo=False, static_one_of_one=True)
    assert footer.paragraphs[0].text.strip() == "pg 1 of 1"
    assert _footer_field_count(doc) == 0


def test_set_report_footer_single_dynamic_page_label():
    doc = Document()
    section = doc.sections[0]
    footer = section.footer
    footer.add_paragraph("page 1 of 1")
    footer.add_paragraph("page 1 of 1")
    set_report_footer(doc, with_logo=True)
    assert len(footer.paragraphs) == 1
    assert footer.paragraphs[0].text.replace(" ", "").startswith("page")
    assert _footer_field_count(doc) == 6
