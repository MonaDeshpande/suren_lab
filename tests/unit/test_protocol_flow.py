"""
Unit tests for analyst protocol flow validation (template routing, page-1
methods, result fill, offline signature blocks).

Maps to manual TC-ANL-020…025 and the Analyst Flow Validation plan.
"""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from services.protocol_docx import (
    JAGGERY_TEMPLATE_PATH,
    NUTRITION_TEMPLATE_PATH,
    OBSERVATION_TABLE_HEADING,
    SUMMARY_HEADER_LABELS,
    WATER_TEMPLATE_PATH,
    _template_path,
    fill_protocol_docx_bytes,
    unsaved_selected_test_names,
)
from services.protocol_store import ProtocolHeader, TestResultRow
from services.protocols.test_catalog import (
    CATEGORY_WATER,
    TEST_CATALOG,
    WATER_MICRO_TEST_KEYS,
    WATER_TEST_KEYS,
    default_test_keys_for_category,
    uses_nutrition_template,
)
from services.samples import SampleRecord

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent


def _sample(**overrides) -> SampleRecord:
    base = dict(
        id=1,
        request_id=1,
        sample_code="SLS-260721-0001",
        sr_no=1,
        sample_name="Test Sample",
        batch_code="",
        quantity="",
        parameters="",
        tests_to_perform="",
        status="pending",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        category="food",
        tests_json="[]",
        lab_code="LAB/CTR/26/001",
        package_type="",
    )
    base.update(overrides)
    return SampleRecord(**base)


def _header() -> ProtocolHeader:
    return ProtocolHeader(
        sample_id=1,
        protocol_no="P-001",
        issued_to="Analyst A",
        issued_by="Reception",
        sample_received_on=date(2026, 7, 21),
        date_of_analysis=date(2026, 7, 21),
        appearance_text="Normal",
    )


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def _cell_has_tc_mar(cell) -> bool:
    tc_pr = cell._tc.tcPr
    return tc_pr is not None and tc_pr.find(qn("w:tcMar")) is not None


def _summary_table_index(doc: Document) -> int | None:
    for idx, table in enumerate(doc.tables):
        if not table.rows:
            continue
        hdr = " ".join(_cell_text(cell) for cell in table.rows[0].cells).lower()
        if "sr" in hdr and "result" in hdr and (
            "parameter" in hdr or "name of test" in hdr
        ):
            return idx
    return None


def _summary_result_by_name(
    doc: Document,
    test_name: str,
    table_idx: int | None = None,
) -> str:
    if table_idx is None:
        table_idx = _summary_table_index(doc)
    if table_idx is None:
        return ""
    table = doc.tables[table_idx]
    for row in table.rows[1:]:
        if test_name.lower() in _cell_text(row.cells[1]).lower():
            return _cell_text(row.cells[3])
    return ""


def _summary_method_by_name(
    doc: Document,
    test_name: str,
    table_idx: int | None = None,
) -> str:
    if table_idx is None:
        table_idx = _summary_table_index(doc)
    if table_idx is None:
        return ""
    table = doc.tables[table_idx]
    for row in table.rows[1:]:
        if test_name.lower() in _cell_text(row.cells[1]).lower():
            return _cell_text(row.cells[2])
    return ""


def _worksheet_table_with_text(doc: Document, needle: str) -> object | None:
    """Find a worksheet table (skip page-1 header and summary tables)."""
    summary_idx = _summary_table_index(doc)
    start = (summary_idx + 1) if summary_idx is not None else 0
    for table in doc.tables[start:]:
        if (
            len(table.rows) == 1
            and len(table.rows[0].cells) == 1
            and "appearance" in _cell_text(table.rows[0].cells[0]).lower()
        ):
            continue
        for row in table.rows:
            for cell in row.cells:
                if needle in _cell_text(cell):
                    return table
    return None


def _template_method_before_fill(template_path: Path, table_idx: int, row_idx: int) -> str:
    doc = Document(str(template_path))
    return _cell_text(doc.tables[table_idx].rows[row_idx].cells[2])


class TestTemplateRouting:
    """TC-ANL-021 / 022 — correct reference template per category + test family."""

    def test_water_uses_water_protocol_template(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        assert _template_path(sample) == WATER_TEMPLATE_PATH
        assert WATER_TEMPLATE_PATH.exists()

    def test_nutrition_food_uses_basic_nutrition_template(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture"]),
            package_type="basic_nutrition",
        )
        assert _template_path(sample) == NUTRITION_TEMPLATE_PATH
        assert NUTRITION_TEMPLATE_PATH.exists()

    def test_fssai_package_uses_jaggery_template_even_if_bn_in_list(self):
        """Package group drives protocol — not stray bn_* on assigned list."""
        sample = _sample(
            tests_json=json.dumps(["moisture", "bn_moisture"]),
            package_type="fssai",
        )
        assert _template_path(sample) == JAGGERY_TEMPLATE_PATH

    def test_jaggery_food_uses_jaggery_template(self):
        sample = _sample(
            tests_json=json.dumps(["moisture"]),
            package_type="fssai",
        )
        assert _template_path(sample) == JAGGERY_TEMPLATE_PATH
        if not JAGGERY_TEMPLATE_PATH.exists():
            pytest.skip("Jaggery template not present in reference/")

    def test_water_auto_assigns_thirteen_tests_at_reception(self):
        """TC-ANL-020 — water bundle: 11 chemical + 2 micro."""
        expected = WATER_TEST_KEYS + WATER_MICRO_TEST_KEYS
        assert default_test_keys_for_category(CATEGORY_WATER) == expected
        assert len(WATER_TEST_KEYS) == 11
        assert len(WATER_MICRO_TEST_KEYS) == 2


@pytest.mark.skipif(not WATER_TEMPLATE_PATH.exists(), reason="Water template missing")
class TestWaterProtocolGeneration:
    """TC-ANL-023 / 024 — page-1 method fixed; result filled from saved data."""

    def test_page1_method_not_overwritten(self):
        template_method = _template_method_before_fill(WATER_TEMPLATE_PATH, 1, 1)
        catalog_method = TEST_CATALOG["ph"].method
        assert template_method == catalog_method

        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        results = [
            TestResultRow(
                test_key="ph",
                test_name="pH",
                method="CUSTOM METHOD OVERRIDE",
                unit="—",
                inputs={"ph_reading": "7.0"},
                result_value="7.0",
                result_numeric=7.0,
            )
        ]
        out = fill_protocol_docx_bytes(
            sample,
            ProtocolHeader(
                sample_id=1,
                protocol_no="P-001",
                issued_to="Analyst A",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 21),
                date_of_analysis=date(2026, 7, 21),
                appearance_text="",
            ),
            results,
            generated_by="Test Analyst",
            generated_at="21/07/2026",
        )
        doc = Document(io.BytesIO(out))
        method_after = _summary_method_by_name(doc, "pH", table_idx=1)
        result_after = _summary_result_by_name(doc, "pH", table_idx=1)

        assert method_after == template_method
        assert method_after != "CUSTOM METHOD OVERRIDE"
        assert result_after == "7.0"

    def test_water_result_table_title_bold_13pt(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        doc = Document(io.BytesIO(fill_protocol_docx_bytes(sample, _header(), [])))
        for paragraph in doc.paragraphs:
            if (paragraph.text or "").strip() == "Result Table:":
                assert paragraph.text == "Result Table:"
                name, size, bold = _paragraph_run_font(paragraph)
                assert name == "Cambria"
                assert size == "26"
                assert bold
                return
        pytest.fail("Result Table: paragraph not found")

    def test_water_summary_data_10pt_header_bold_11pt(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        results = [
            TestResultRow(
                test_key="ph",
                test_name="pH",
                method=TEST_CATALOG["ph"].method,
                unit="—",
                inputs={},
                result_value="6.8",
                result_numeric=6.8,
            )
        ]
        doc = Document(io.BytesIO(fill_protocol_docx_bytes(sample, _header(), results)))
        summary = doc.tables[1]
        for cell in summary.rows[0].cells:
            name, size, bold = _cell_run_font(cell)
            assert name == "Cambria"
            assert size == "22"
            assert bold
        for row in summary.rows[1:]:
            for cell in row.cells:
                name, size, bold = _cell_run_font(cell)
                if name is None and size is None:
                    continue
                assert name == "Cambria"
                assert size == "20"
                assert not bold

    def test_water_protocol_header_row_10pt(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        doc = Document(io.BytesIO(fill_protocol_docx_bytes(sample, _header(), [])))
        for cell in doc.tables[0].rows[0].cells:
            name, size, bold = _cell_run_font(cell)
            if name is None and size is None:
                continue
            assert name == "Cambria"
            assert size == "20"
            assert not bold
            assert _cell_has_no_wrap(cell)

    def test_water_blank_line_before_result_table(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        doc = Document(io.BytesIO(fill_protocol_docx_bytes(sample, _header(), [])))
        assert _has_blank_line_before_result_table(doc)

    def test_page1_result_filled_for_saved_tests_only(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(WATER_TEST_KEYS))
        header = ProtocolHeader(
            sample_id=1,
            protocol_no="P-001",
            issued_to="Analyst A",
            issued_by="Reception",
            sample_received_on=date(2026, 7, 21),
            date_of_analysis=date(2026, 7, 21),
            appearance_text="",
        )
        results = [
            TestResultRow(
                test_key="ph",
                test_name="pH",
                method=TEST_CATALOG["ph"].method,
                unit="—",
                inputs={},
                result_value="6.8",
                result_numeric=6.8,
            )
        ]
        out = fill_protocol_docx_bytes(sample, header, results)
        doc = Document(io.BytesIO(out))
        t1 = doc.tables[1]
        filled = [
            _cell_text(row.cells[3])
            for row in t1.rows[1:]
            if _cell_text(row.cells[3])
        ]
        assert filled == ["6.8"]
        assert len(t1.rows) == 12  # header + 11 fixed water parameters

    def test_page1_keeps_all_eleven_template_methods(self):
        template_doc = Document(str(WATER_TEMPLATE_PATH))
        template_methods = [
            _cell_text(template_doc.tables[1].rows[i].cells[2])
            for i in range(1, 12)
        ]
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        out = fill_protocol_docx_bytes(
            sample,
            _header(),
            [
                TestResultRow(
                    test_key="ph",
                    test_name="pH",
                    method="OVERRIDE",
                    unit="",
                    inputs={},
                    result_value="7.0",
                    result_numeric=7.0,
                )
            ],
        )
        doc = Document(io.BytesIO(out))
        filled_methods = [
            _cell_text(doc.tables[1].rows[i].cells[2]) for i in range(1, 12)
        ]
        assert filled_methods == template_methods

    def test_water_micro_observation_table_on_last_page(self):
        from services.protocols.test_catalog import default_water_micro_procedure

        sample = _sample(
            category=CATEGORY_WATER,
            tests_json=json.dumps(WATER_TEST_KEYS + WATER_MICRO_TEST_KEYS),
        )
        results = [
            TestResultRow(
                test_key="water_total_coliform",
                test_name="Total Coliform",
                method="",
                unit="",
                inputs={
                    "procedure": default_water_micro_procedure("water_total_coliform"),
                    "result_obs": "present",
                },
                result_value="present",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="water_e_coli",
                test_name="E. coli",
                method="",
                unit="",
                inputs={
                    "procedure": "Custom E. coli procedure",
                    "result_obs": "absent",
                },
                result_value="absent",
                result_numeric=None,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        full_text = _full_document_text(doc)
        assert "OBSERVATION TABLE" in full_text
        assert "Total Coliform" in full_text
        assert "E. coli" in full_text
        assert "present" in full_text
        assert "Custom E. coli procedure" in full_text
        assert "absent" in full_text
        # Page-1 summary stays chemical-only (11 data rows)
        assert len(doc.tables[1].rows) == 12


@pytest.mark.skipif(not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing")
class TestNutritionProtocolGeneration:
    """TC-ANL-022 — Basic Nutrition template + fixed methods."""

    def test_page1_method_not_overwritten(self):
        template_method = _template_method_before_fill(NUTRITION_TEMPLATE_PATH, 1, 2)
        catalog_method = TEST_CATALOG["bn_moisture"].method
        assert template_method == catalog_method

        sample = _sample(tests_json=json.dumps(["bn_moisture"]))
        results = [
            TestResultRow(
                test_key="bn_moisture",
                test_name="Moisture",
                method="OVERRIDE",
                unit="%",
                inputs={},
                result_value="5.0",
                result_numeric=5.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        method_after = _summary_method_by_name(doc, "Moisture", table_idx=1)
        assert method_after == template_method
        assert _summary_result_by_name(doc, "Moisture", table_idx=1) == "5.0"


@pytest.mark.skipif(not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing")
class TestNutritionWorksheetReadings:
    """Basic Nutrition observation tables: readings + worked = answer, format-safe."""

    def _nutrition_moisture_results(self) -> list[TestResultRow]:
        return [
            TestResultRow(
                test_key="bn_moisture",
                test_name="Moisture",
                method=TEST_CATALOG["bn_moisture"].method,
                unit="%",
                inputs={
                    "empty_dish": 68.7246,
                    "w1": 73.7923,
                    "w": 5.0677,
                    "after_dry": 73.6292,
                    "w2": 73.6292,
                },
                result_value="3.21",
                result_numeric=3.21,
            )
        ]

    def test_bn_moisture_readings_and_worked_formula(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture"]),
            package_type="basic_nutrition",
        )
        out = fill_protocol_docx_bytes(
            sample, _header(), self._nutrition_moisture_results()
        )
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        assert "68.7246" in _cell_text(table.rows[1].cells[1])
        assert "73.7923" in _cell_text(table.rows[2].cells[1])
        assert "5.0677" in _cell_text(table.rows[3].cells[1])
        assert "73.6292" in _cell_text(table.rows[5].cells[1])
        formula_row = table.rows[7]
        readings = _cell_text(formula_row.cells[1])
        assert "=" in readings
        assert "3.21" in readings
        desc = _cell_text(formula_row.cells[0])
        assert "73.7923" not in desc
        assert "Moisture" in desc or "W1" in desc

    def test_bn_moisture_preserves_template_formula_wording(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture"]),
            package_type="basic_nutrition",
        )
        out = fill_protocol_docx_bytes(
            sample, _header(), self._nutrition_moisture_results()
        )
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        filled_desc = _cell_text(table.rows[7].cells[0])
        assert "Moisture w/w" in filled_desc or "W1" in filled_desc
        assert "Moisture % =" not in filled_desc

    def test_bn_moisture_formula_readings_alignment_preserved(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        ref = Document(str(NUTRITION_TEMPLATE_PATH))
        ref_align = ref.tables[3].rows[7].cells[1].paragraphs[0].alignment
        sample = _sample(
            tests_json=json.dumps(["bn_moisture"]),
            package_type="basic_nutrition",
        )
        out = fill_protocol_docx_bytes(
            sample, _header(), self._nutrition_moisture_results()
        )
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        filled_align = table.rows[7].cells[1].paragraphs[0].alignment
        assert filled_align == ref_align == WD_ALIGN_PARAGRAPH.RIGHT

    def test_bn_ash_insoluble_dual_worked_lines(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture", "bn_ash_insoluble_hcl"]),
            package_type="basic_nutrition",
        )
        results = [
            *self._nutrition_moisture_results(),
            TestResultRow(
                test_key="bn_ash_insoluble_hcl",
                test_name="Ash Insoluble in HCL",
                method=TEST_CATALOG["bn_ash_insoluble_hcl"].method,
                unit="%",
                inputs={
                    "w1": 45.0958,
                    "w": 5.0854,
                    "w2": 45.0970,
                },
                result_value="0.023",
                result_numeric=0.023,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "Insoluble ash")
        assert table is not None
        wet_row = table.rows[6]
        dry_row = table.rows[7]
        assert "45.097" in _cell_text(wet_row.cells[1])
        assert "0.023" in _cell_text(dry_row.cells[1])
        assert "96.79" in _cell_text(dry_row.cells[1]) or "3.21" in _cell_text(
            dry_row.cells[1]
        )

    def test_bn_protein_paragraph_worked_lines(self):
        sample = _sample(
            tests_json=json.dumps(["bn_protein"]),
            package_type="basic_nutrition",
        )
        results = [
            TestResultRow(
                test_key="bn_protein",
                test_name="Protein",
                method=TEST_CATALOG["bn_protein"].method,
                unit="%",
                inputs={
                    "w": 1.5,
                    "n_naoh": 0.1,
                    "br_blank": 10.0,
                    "br_sample": 8.5,
                    "n_factor": 6.25,
                },
                result_value="6.5",
                result_numeric=6.5,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        body = "\n".join(p.text for p in doc.paragraphs)
        assert "0.014" in body
        assert "6.5" in body
        assert "× 6.25" in body or "x 6.25" in body.lower()

    def test_bn_sugar_readings_and_formula(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture", "bn_added_sugar", "bn_total_sugar"]),
            package_type="basic_nutrition",
        )
        results = [
            *self._nutrition_moisture_results(),
            TestResultRow(
                test_key="bn_added_sugar",
                test_name="Added Sugar",
                method=TEST_CATALOG["bn_added_sugar"].method,
                unit="%",
                inputs={
                    "sample_wt": 1.0201,
                    "fehling": 10,
                    "br": 23.5,
                    "sugar_conc": 0.21775,
                },
                result_value="14.84",
                result_numeric=14.84,
            ),
            TestResultRow(
                test_key="bn_total_sugar",
                test_name="Total Sugar",
                method=TEST_CATALOG["bn_total_sugar"].method,
                unit="%",
                inputs={
                    "sample_wt": 1.0201,
                    "fehling": 10,
                    "br": 22.5,
                    "sugar_conc": 0.227,
                },
                result_value="91.96",
                result_numeric=91.96,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "Fehling solution")
        assert table is not None
        assert "1.0201" in _cell_text(table.rows[1].cells[2])
        assert "23.5" in _cell_text(table.rows[3].cells[2])
        assert "22.5" in _cell_text(table.rows[5].cells[2])
        formula_cell = table.rows[6].cells[0]
        formula_text = _cell_text(formula_cell)
        assert "Concentration of sugar" in formula_text
        assert "14.84" in formula_text
        assert "91.96" in formula_text
        assert "Total Invert sugar" in formula_text
        assert "Total reducing sugar" in formula_text

    def test_nutrition_fill_preview_docx_written(self, tmp_path):
        sample = _sample(
            sample_name="Protein Sauce",
            lab_code="GLG/26/555",
            tests_json=json.dumps(["bn_moisture", "bn_total_ash"]),
            package_type="basic_nutrition",
        )
        results = [
            *self._nutrition_moisture_results(),
            TestResultRow(
                test_key="bn_total_ash",
                test_name="Total Ash",
                method=TEST_CATALOG["bn_total_ash"].method,
                unit="%",
                inputs={"w1": 45.0958, "w": 5.0854, "w2": 45.1114},
                result_value="0.30",
                result_numeric=0.30,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        preview = PROJECT_ROOT / "tmp" / "nutrition_fill_preview.docx"
        preview.parent.mkdir(parents=True, exist_ok=True)
        preview.write_bytes(out)
        assert preview.exists()
        assert preview.stat().st_size > 1000


class TestHeaderValueColumns:
    """Sample name / dates go in the empty value cell beside each label."""

    def _assert_label_value(self, label: str, value: str, label_cell, value_cell):
        assert label in _cell_text(label_cell)
        assert value == _cell_text(value_cell)
        assert value not in _cell_text(label_cell) or _cell_text(label_cell) == label

    @pytest.mark.skipif(not JAGGERY_TEMPLATE_PATH.exists(), reason="Jaggery template missing")
    def test_jaggery_header_values_in_next_column(self):
        sample = _sample(
            sample_name="Jaggery Lot A",
            lab_code="LAB/CTR/26/099",
            tests_json=json.dumps(["moisture"]),
        )
        header = _header()
        out = fill_protocol_docx_bytes(sample, header, [])
        doc = Document(io.BytesIO(out))
        t0 = doc.tables[0]
        assert _cell_text(t0.rows[0].cells[0]).startswith("Sample Name")
        assert _cell_text(t0.rows[0].cells[1]) == "Jaggery Lot A"
        assert "Jaggery Lot A" not in _cell_text(t0.rows[0].cells[0])
        assert _cell_text(t0.rows[0].cells[2]).startswith("Sample Received on")
        assert _cell_text(t0.rows[0].cells[3]) == "21/07/2026"
        assert _cell_text(t0.rows[1].cells[1]) == "LAB/CTR/26/099"
        assert _cell_text(t0.rows[1].cells[3]) == "21/07/2026"

    @pytest.mark.skipif(not WATER_TEMPLATE_PATH.exists(), reason="Water template missing")
    def test_water_header_values_in_next_column(self):
        sample = _sample(
            sample_name="Borewell Water",
            lab_code="LAB/CTR/26/088",
            category=CATEGORY_WATER,
            tests_json=json.dumps(["ph"]),
        )
        out = fill_protocol_docx_bytes(sample, _header(), [])
        doc = Document(io.BytesIO(out))
        t0 = doc.tables[0]
        assert _cell_text(t0.rows[0].cells[1]) == "P-001"
        assert _cell_text(t0.rows[2].cells[3]) == "Borewell Water"
        assert "Borewell Water" not in _cell_text(t0.rows[2].cells[0])
        assert _cell_text(t0.rows[2].cells[7]) == "21/07/2026"
        assert _cell_text(t0.rows[3].cells[3]) == "LAB/CTR/26/088"
        assert _cell_text(t0.rows[3].cells[7]) == "21/07/2026"

    @pytest.mark.skipif(not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing")
    def test_nutrition_header_values_in_next_column(self):
        sample = _sample(
            sample_name="Protein Mix",
            lab_code="LAB/CTR/26/077",
            tests_json=json.dumps(["bn_moisture"]),
        )
        out = fill_protocol_docx_bytes(sample, _header(), [])
        doc = Document(io.BytesIO(out))
        t0 = doc.tables[0]
        assert _cell_text(t0.rows[0].cells[1]) == "Protein Mix"
        assert "Protein Mix" not in _cell_text(t0.rows[0].cells[0])
        assert _cell_text(t0.rows[0].cells[3]) == "21/07/2026"
        assert _cell_text(t0.rows[1].cells[1]) == "LAB/CTR/26/077"
        assert _cell_text(t0.rows[1].cells[3]) == "21/07/2026"
        hdr = doc.sections[0].header.tables[0]
        assert _cell_text(hdr.rows[0].cells[1]) == "P-001"
        assert _cell_text(hdr.rows[0].cells[3]) == "Analyst A"
        assert _cell_text(hdr.rows[0].cells[5]) == "Reception"


def _footer_text(doc: Document) -> str:
    text = ""
    for section in doc.sections:
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += "\n" + _cell_text(cell)
        for para in section.footer.paragraphs:
            text += "\n" + (para.text or "")
    return text


def _full_document_text(doc: Document) -> str:
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + _cell_text(cell)
    text += _footer_text(doc)
    return text


class TestOfflineSignature:
    """TC-ANL-025 — Basic Nutrition-style footer; no body signature blocks."""

    def _assert_standard_protocol_footer(self, doc: Document) -> None:
        footer = _footer_text(doc)
        assert "Prepared by" in footer
        assert "Reviewed" in footer and "Issued" in footer
        assert "Approved by" in footer
        assert "Surendra" not in footer
        assert "Nashikkar" not in footer
        assert "Control copy" not in footer

    def _assert_no_body_signatures(self, doc: Document) -> None:
        body_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Analyzed By" not in body_text
        assert "Checked By" not in body_text
        assert "Technical Manager" not in body_text
        assert "Generated by" not in body_text

    @pytest.mark.skipif(not WATER_TEMPLATE_PATH.exists(), reason="Water template missing")
    def test_water_protocol_footer_and_no_body_signatures(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        out = fill_protocol_docx_bytes(
            sample,
            _header(),
            [],
            generated_by="Analyst One",
            generated_at="21/07/2026 14:00",
        )
        doc = Document(io.BytesIO(out))
        footer = _footer_text(doc)
        assert "Prepared" in footer and "Issued" in footer
        assert "Surendra" not in footer
        assert "Nashikkar" not in footer
        self._assert_no_body_signatures(doc)

    @pytest.mark.skipif(not JAGGERY_TEMPLATE_PATH.exists(), reason="Jaggery template missing")
    def test_jaggery_footer_matches_nutrition_pattern(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        self._assert_standard_protocol_footer(doc)
        self._assert_no_body_signatures(doc)
        assert len(doc.sections) == 1
        full_text = _full_document_text(doc)
        assert "Generated by" not in full_text


@pytest.mark.skipif(not JAGGERY_TEMPLATE_PATH.exists(), reason="Jaggery template missing")
class TestJaggeryWorksheetReadings:
    """Saved Moisture inputs and answer must appear in the observation table."""

    def test_moisture_readings_and_answer_filled(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={
                    "empty_dish": 50.0,
                    "w1": 55.0,
                    "w": 5.0,
                    "after_dry": 54.5,
                    "w2": 54.5,
                },
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        t5 = _worksheet_table_with_text(doc, "stainless steel dish")
        assert t5 is not None
        assert "50" in _cell_text(t5.rows[1].cells[1])
        assert "55" in _cell_text(t5.rows[2].cells[1])
        assert "5" in _cell_text(t5.rows[3].cells[1])
        assert "54.5" in _cell_text(t5.rows[4].cells[1])
        assert "54.5" in _cell_text(t5.rows[5].cells[1])
        # Formula answer in Readings/% column
        assert "10" in _cell_text(t5.rows[6].cells[1])

    def test_moisture_symbolic_in_description_calculation_in_readings(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        formula_row = table.rows[6]
        desc = _cell_text(formula_row.cells[0])
        readings = _cell_text(formula_row.cells[1])
        assert "55" not in desc
        assert "54.5" not in desc
        assert "Moisture" in desc or "W1" in desc
        assert "=" in readings
        assert "10" in readings
        assert "%" in readings

    def test_total_ash_dual_worked_lines(self):
        sample = _sample(tests_json=json.dumps(["moisture", "total_ash"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.8235},
                result_value="3.53",
                result_numeric=3.53,
            ),
            TestResultRow(
                test_key="total_ash",
                test_name=TEST_CATALOG["total_ash"].name,
                method=TEST_CATALOG["total_ash"].method,
                unit="%",
                inputs={
                    "w1": 45.0,
                    "before_ign": 50.0,
                    "w": 5.0,
                    "after_ign": 45.05895,
                    "w2": 45.05895,
                },
                result_value="1.22",
                result_numeric=1.22,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "Ash w/w")
        assert table is not None
        wet_row = table.rows[6]
        dry_row = table.rows[7]
        wet_read = _cell_text(wet_row.cells[1])
        dry_read = _cell_text(dry_row.cells[1])
        assert "=" in wet_read
        assert "45" in wet_read
        assert "%" in wet_read
        assert "=" in dry_read
        assert "1.22" in dry_read
        assert "3.53" in dry_read

    def test_worksheet_cell_padding_applied(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        for row in table.rows:
            for cell in row.cells:
                assert _cell_has_tc_mar(cell)

    def test_extraneous_dual_formula_rows_split_columns(self):
        sample = _sample(tests_json=json.dumps(["extraneous_matter"]))
        results = [
            TestResultRow(
                test_key="extraneous_matter",
                test_name=TEST_CATALOG["extraneous_matter"].name,
                method=TEST_CATALOG["extraneous_matter"].method,
                unit="%",
                inputs={
                    "w": 2.0045,
                    "w2_filter": 1.0261,
                    "w1_matter": 2.01,
                    "moisture_pct": 3.22,
                },
                result_value="50.72",
                result_numeric=50.72,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "Filter Paper")
        assert table is not None
        wet_row = table.rows[4]
        dry_row = table.rows[5]
        wet_desc = _cell_text(wet_row.cells[0])
        wet_read = _cell_text(wet_row.cells[1])
        dry_desc = _cell_text(dry_row.cells[0])
        dry_read = _cell_text(dry_row.cells[1])
        assert "2.01" not in wet_desc
        assert "Extraneous matter" in wet_desc
        assert "2.01" in wet_read
        assert "49.08" in wet_read
        assert "dry basis" in dry_desc.lower()
        assert "50.72" in dry_read

    def test_page1_reducing_sugar_not_stolen_by_nutrition_name(self):
        """Fuzzy catalog match used to map reducing sugar → bn_total_sugar."""
        sample = _sample(
            tests_json=json.dumps(["invert_sugar", "reducing_sugar", "sucrose"])
        )
        results = [
            TestResultRow(
                test_key="invert_sugar",
                test_name=TEST_CATALOG["invert_sugar"].name,
                method=TEST_CATALOG["invert_sugar"].method,
                unit="%",
                inputs={
                    "sample_wt": 1.0201,
                    "fehling_invert": 10,
                    "br_invert": 22.5,
                    "sugar_conc": 0.227,
                    "moisture_pct": 3.21,
                },
                result_value="91.96",
                result_numeric=91.96,
            ),
            TestResultRow(
                test_key="reducing_sugar",
                test_name=TEST_CATALOG["reducing_sugar"].name,
                method=TEST_CATALOG["reducing_sugar"].method,
                unit="%",
                inputs={
                    "sample_wt": 1.5148,
                    "fehling_reducing": 10,
                    "br_reducing": 23.5,
                    "sugar_conc": 0.21775,
                    "moisture_pct": 3.21,
                },
                result_value="14.84",
                result_numeric=14.84,
            ),
            TestResultRow(
                test_key="sucrose",
                test_name=TEST_CATALOG["sucrose"].name,
                method=TEST_CATALOG["sucrose"].method,
                unit="%",
                inputs={},
                result_value="73.26",
                result_numeric=73.26,
            ),
        ]
        out = fill_protocol_docx_bytes(
            sample, _header(), results, generated_by="S. Ambilwade"
        )
        doc = Document(io.BytesIO(out))
        assert _summary_result_by_name(doc, "Invert Sugar") == "91.96"
        assert _summary_result_by_name(doc, "Reduced sugar") == "14.84"
        assert _summary_result_by_name(doc, "Sucrose") == "73.26"
        sugar = _worksheet_table_with_text(doc, "Fehling solution")
        assert sugar is not None
        sugar_formula_row = sugar.rows[6]
        cell_text = _cell_text(sugar_formula_row.cells[0])
        assert "Total Invert sugar" in cell_text or "Concentration of sugar" in cell_text
        assert "Wt. of sample taken in g X B.R." in cell_text
        assert "91.96" in cell_text
        assert "14.84" in cell_text
        assert "73.26" in cell_text
        assert "× 250" in cell_text
        sugar_text = "\n".join(
            _cell_text(cell)
            for row in sugar.rows
            for cell in row.cells
        )
        assert "1.5148" in sugar_text
        body = "\n".join(p.text for p in doc.paragraphs)
        assert "Analyzed By" not in body
        assert "Technical Manager" not in body

    def test_sulphated_ash_and_so2_worksheets_filled(self):
        sample = _sample(tests_json=json.dumps(["sulphated_ash", "sulphur_dioxide"]))
        results = [
            TestResultRow(
                test_key="sulphated_ash",
                test_name=TEST_CATALOG["sulphated_ash"].name,
                method=TEST_CATALOG["sulphated_ash"].method,
                unit="%",
                inputs={
                    "w1": 45.0958,
                    "before_ign": 50.1812,
                    "w": 5.0854,
                    "after_ign": 45.1114,
                    "w2": 45.1114,
                    "moisture_pct": 3.21,
                },
                result_value="0.30",
                result_numeric=0.30,
            ),
            TestResultRow(
                test_key="sulphur_dioxide",
                test_name=TEST_CATALOG["sulphur_dioxide"].name,
                method=TEST_CATALOG["sulphur_dioxide"].method,
                unit="ppm",
                inputs={"sample_wt": 25.0180, "ug_so4": 16.94},
                result_value="6.77",
                result_numeric=6.77,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        ash = _worksheet_table_with_text(doc, "empty Crucible")
        so2 = _worksheet_table_with_text(doc, "µg SO")
        assert ash is not None and so2 is not None
        assert "45.0958" in _cell_text(ash.rows[1].cells[1])
        assert "5.0854" in _cell_text(ash.rows[3].cells[1])
        assert "0.30" in _cell_text(ash.rows[8].cells[1])
        assert "25.018" in _cell_text(so2.rows[1].cells[1])
        assert "6.77" in _cell_text(so2.rows[2].cells[1])
        assert _summary_result_by_name(doc, "Sulphated ash") == "0.30"
        assert _summary_result_by_name(doc, "Sulphur dioxide") == "6.77"

    def test_unsaved_selected_tests_listed(self):
        sample = _sample(tests_json=json.dumps(["moisture", "total_ash"]))
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        missing = unsaved_selected_test_names(sample, results)
        assert any("ash" in n.lower() for n in missing)
        assert not any("moisture" in n.lower() for n in missing)


@pytest.mark.skipif(not JAGGERY_TEMPLATE_PATH.exists(), reason="Jaggery template missing")
class TestJaggeryPruneAndResultCells:
    """Unconducted tests must not leave orphan headings; choice results fill worksheets."""

    def test_unconducted_tests_remove_section_headings(self):
        sample = _sample(
            sample_name="Masala",
            tests_json=json.dumps(["moisture", "added_color"]),
        )
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 86.0, "w": 6.0, "w2": 82.0},
                result_value="66.67",
                result_numeric=66.67,
            ),
            TestResultRow(
                test_key="added_color",
                test_name="Added Color",
                method=TEST_CATALOG["added_color"].method,
                unit="",
                inputs={"color_result": "Absent"},
                result_value="Absent",
                result_numeric=None,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        body_text = "\n".join(p.text for p in doc.paragraphs)
        assert "TOTAL ASH (on dry basis)" not in body_text
        assert "ASH INSOLUBLE" not in body_text

    def test_moisture_and_added_color_are_separate_sections(self):
        sample = _sample(
            sample_name="Masala",
            tests_json=json.dumps(["moisture", "added_color"]),
        )
        results = [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 86.0, "w": 6.0, "w2": 82.0},
                result_value="66.67",
                result_numeric=66.67,
            ),
            TestResultRow(
                test_key="added_color",
                test_name="Added Color",
                method=TEST_CATALOG["added_color"].method,
                unit="",
                inputs={"color_result": "Absent"},
                result_value="Absent",
                result_numeric=None,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))

        section_titles: list[str] = []
        seen_obs = False
        for child in doc.element.body:
            if child.tag.endswith("p"):
                text = next(
                    (p.text or "").strip()
                    for p in doc.paragraphs
                    if p._element is child
                )
                if not text:
                    continue
                if "observation table" in text.lower():
                    seen_obs = True
                    continue
                if not seen_obs or text == "Result Table:":
                    continue
                first_line = text.split("\n", 1)[0].strip()
                if first_line.endswith(":"):
                    section_titles.append(first_line)
            elif child.tag.endswith("tbl") and seen_obs:
                table = next(t for t in doc.tables if t._tbl is child)
                if (
                    len(table.rows) == 1
                    and len(table.rows[0].cells) == 1
                    and "appearance" in _cell_text(table.rows[0].cells[0]).lower()
                ):
                    continue

        moisture_titles = [t for t in section_titles if "MOISTURE" in t.upper()]
        color_titles = [t for t in section_titles if "ADDED COLOR" in t.upper()]
        assert moisture_titles, "Expected MOISTURE section title"
        assert color_titles, "Expected ADDED COLOR section title"
        assert moisture_titles[0].startswith("1.")
        assert color_titles[0].startswith("2.")

        moisture = _worksheet_table_with_text(doc, "stainless steel dish")
        color = _worksheet_table_with_text(doc, "Added Color")
        assert moisture is not None
        assert color is not None
        assert moisture is not color

        assert _cell_text(color.rows[0].cells[1]).strip() == "Result"
        assert "Absent" in _cell_text(color.rows[1].cells[1])

        prev = color._tbl.getprevious()
        assert prev is not None and prev.tag.endswith("p")
        para = Paragraph(prev, doc)
        assert "ADDED COLOR" in (para.text or "").upper()
        assert _paragraph_has_keep_next(para)

    def test_added_color_worksheet_result_filled(self):
        sample = _sample(tests_json=json.dumps(["added_color"]))
        results = [
            TestResultRow(
                test_key="added_color",
                test_name="Added Color",
                method=TEST_CATALOG["added_color"].method,
                unit="",
                inputs={"color_result": "Absent"},
                result_value="Absent",
                result_numeric=None,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "Added Color")
        assert table is not None
        assert "Absent" in _cell_text(table.rows[1].cells[1])
        assert _cell_text(table.rows[0].cells[1]).strip() == "Result"
        assert _summary_result_by_name(doc, "Added Color") == "Absent"


@pytest.mark.skipif(
    not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing"
)
class TestDirectorApprovalBlanked:
    """Approved-by footer must not keep Surendra name/signature."""

    def test_nutrition_footer_surendra_removed(self):
        sample = _sample(tests_json=json.dumps(["bn_moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), [])
        doc = Document(io.BytesIO(out))
        footer_text = ""
        for section in doc.sections:
            for table in section.footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        footer_text += "\n" + _cell_text(cell)
            for para in section.footer.paragraphs:
                footer_text += "\n" + para.text
        assert "Surendra" not in footer_text
        assert "Nashikkar" not in footer_text
        assert "Approved by" in footer_text


def _document_has_page_break(doc: Document) -> bool:
    for child in doc.element.body:
        if not child.tag.endswith("p"):
            continue
        for br in child.findall(".//" + qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True
    return False


def _observation_heading_index(doc: Document) -> int | None:
    for idx, paragraph in enumerate(doc.paragraphs):
        if "observation table" in (paragraph.text or "").lower():
            return idx
    return None


def _page_break_before_observation(doc: Document) -> bool:
    obs_idx = _observation_heading_index(doc)
    if obs_idx is None or obs_idx == 0:
        return False
    prev = doc.paragraphs[obs_idx - 1]._element
    for br in prev.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _body_index_after_summary(doc: Document) -> int | None:
    summary_idx = _summary_table_index(doc)
    if summary_idx is None:
        return None
    summary_tbl = doc.tables[summary_idx]._tbl
    siblings = list(doc.element.body)
    try:
        return siblings.index(summary_tbl) + 1
    except ValueError:
        return None


def _first_worksheet_body_index(doc: Document) -> int | None:
    start = _body_index_after_summary(doc)
    if start is None:
        return None
    siblings = list(doc.element.body)
    for idx in range(start, len(siblings)):
        child = siblings[idx]
        if not child.tag.endswith(("p", "tbl")):
            continue
        if child.tag.endswith("p"):
            text = next(
                (p.text or "").strip() for p in doc.paragraphs if p._element is child
            )
            if not text or text == "Result Table:":
                continue
            if OBSERVATION_TABLE_HEADING.lower() in text.lower():
                continue
            if text.endswith(":"):
                return idx
        elif child.tag.endswith("tbl"):
            table = next(t for t in doc.tables if t._tbl is child)
            hdr = (table.rows[0].cells[0].text or "").lower()
            if "description" in hdr or "appearance" in hdr:
                return idx
    return None


@pytest.mark.skipif(not JAGGERY_TEMPLATE_PATH.exists(), reason="Jaggery template missing")
def _table_grid_col_inches(table) -> list[float]:
    return [tw / 1440 for tw in _table_grid_col_twips(table)]


def _table_grid_col_twips(table) -> list[int]:
    from docx.oxml.ns import qn

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    return [int(col.get(qn("w:w"))) for col in grid.findall(qn("w:gridCol"))]


def _usable_page_width_inches(section) -> float:
    return section.page_width.inches - section.left_margin.inches - section.right_margin.inches


def _table_tbl_ind(table) -> tuple[str, str] | None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return None
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        return None
    return (ind.get(qn("w:w")), ind.get(qn("w:type")))


def _table_grid_sum_twips(table) -> int:
    return sum(_table_grid_col_twips(table))


def _table_outer_right_edge_twips(table) -> int:
    ind = _table_tbl_ind(table)
    ind_twips = int(ind[0]) if ind is not None else 0
    return ind_twips + _table_grid_sum_twips(table)


def _canonical_worksheet_outer_frame() -> tuple[int, tuple[str, str]]:
    ref = Document(str(NUTRITION_TEMPLATE_PATH))
    worksheet_ref = ref.tables[3]
    target_width = _table_grid_sum_twips(worksheet_ref)
    target_ind = _table_tbl_ind(worksheet_ref) or ("-162", "dxa")
    return target_width, target_ind


def _empty_paragraphs_between_summary_and_observation(doc: Document) -> int:
    summary_idx = _summary_table_index(doc)
    if summary_idx is None:
        return 0
    summary_tbl = doc.tables[summary_idx]._tbl
    count = 0
    next_el = summary_tbl.getnext()
    while next_el is not None:
        if next_el.tag.endswith("p"):
            para = Paragraph(next_el, doc)
            text = (para.text or "").strip()
            if "observation table" in text.lower():
                break
            if not text:
                count += 1
        elif next_el.tag.endswith("tbl"):
            break
        next_el = next_el.getnext()
    return count


class TestProtocolPageLayout:
    """Page 1 = result summary only; page 2+ = observation worksheets."""

    def _jaggery_results(self) -> list[TestResultRow]:
        return [
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]

    def test_jaggery_has_observation_heading_and_page_break(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        assert _observation_heading_index(doc) is not None
        assert _document_has_page_break(doc)
        assert _page_break_before_observation(doc)

    def test_jaggery_summary_headers_match_reference(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        summary_idx = _summary_table_index(doc)
        assert summary_idx is not None
        headers = [_cell_text(cell) for cell in doc.tables[summary_idx].rows[0].cells]
        assert headers[:5] == list(SUMMARY_HEADER_LABELS)

    def test_jaggery_margins_match_reference(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        section = doc.sections[0]
        assert section.left_margin == Inches(1)
        assert section.right_margin == Inches(1)
        assert section.top_margin == Inches(1)
        assert section.bottom_margin == Inches(1)

    def test_jaggery_worksheet_starts_after_summary_block(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        summary_body_idx = _body_index_after_summary(doc)
        worksheet_body_idx = _first_worksheet_body_index(doc)
        obs_body_idx = next(
            (
                idx
                for idx, child in enumerate(doc.element.body)
                if child.tag.endswith("p")
                and "observation table"
                in next(
                    (p.text or "").lower()
                    for p in doc.paragraphs
                    if p._element is child
                )
            ),
            None,
        )
        assert summary_body_idx is not None
        assert worksheet_body_idx is not None
        assert obs_body_idx is not None
        assert summary_body_idx < obs_body_idx < worksheet_body_idx

    @pytest.mark.skipif(
        not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing"
    )
    def test_nutrition_page_break_before_observation(self):
        sample = _sample(
            tests_json=json.dumps(["bn_moisture"]),
            package_type="basic_nutrition",
        )
        results = [
            TestResultRow(
                test_key="bn_moisture",
                test_name="Moisture",
                method=TEST_CATALOG["bn_moisture"].method,
                unit="%",
                inputs={"w1": 55.0, "w": 5.0, "w2": 54.5},
                result_value="10.0",
                result_numeric=10.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        assert _observation_heading_index(doc) is not None
        assert _page_break_before_observation(doc)

    @pytest.mark.skipif(not WATER_TEMPLATE_PATH.exists(), reason="Water template missing")
    def test_water_page_break_and_footer(self):
        sample = _sample(category=CATEGORY_WATER, tests_json=json.dumps(["ph"]))
        results = [
            TestResultRow(
                test_key="ph",
                test_name="pH",
                method=TEST_CATALOG["ph"].method,
                unit="—",
                inputs={"ph_reading": "7.0"},
                result_value="7.0",
                result_numeric=7.0,
            )
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        assert _observation_heading_index(doc) is not None
        assert _page_break_before_observation(doc)
        footer = _footer_text(doc)
        assert "Surendra" not in footer
        assert "Nashikkar" not in footer
        ref_sec = Document(str(WATER_TEMPLATE_PATH)).sections[0]
        section = doc.sections[0]
        assert section.left_margin == ref_sec.left_margin
        assert section.right_margin == ref_sec.right_margin

    def test_jaggery_reference_page1_sample_table(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        canonical_width, canonical_ind = _canonical_worksheet_outer_frame()
        t0 = doc.tables[0]
        assert len(t0.rows) == 2
        assert _cell_text(t0.rows[0].cells[0]).startswith("Sample Name")
        assert _cell_text(doc.tables[1].rows[0].cells[0]).startswith("Sr")
        assert _table_tbl_ind(t0) == canonical_ind
        assert _table_grid_sum_twips(t0) == canonical_width
        tbl_pr = t0._tbl.tblPr
        assert tbl_pr is None or tbl_pr.find(qn("w:tblpPr")) is None

    def test_repeating_section_header_table(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        hdr_tables = doc.sections[0].header.tables
        assert hdr_tables
        hdr_text = _cell_text(hdr_tables[0].rows[0].cells[0])
        assert "Protocol No" in hdr_text
        assert _cell_text(hdr_tables[0].rows[0].cells[1]) == "P-001"

    def test_appearance_worksheet_filled(self):
        sample = _sample(tests_json=json.dumps(["appearance", "moisture"]))
        results = [
            TestResultRow(
                test_key="appearance",
                test_name="Appearance",
                method="",
                unit="",
                inputs={},
                result_value="Red Colored Viscous",
                result_numeric=None,
            ),
            *self._jaggery_results(),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        app_table = _appearance_table(doc)
        assert app_table is not None
        assert "Red Colored Viscous" in _cell_text(app_table.rows[0].cells[0])

    def test_bare_unit_placeholders_cleared(self):
        sample = _sample(tests_json=json.dumps(["moisture"]))
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        table = _worksheet_table_with_text(doc, "stainless steel dish")
        assert table is not None
        for row in table.rows[1:]:
            desc = _cell_text(row.cells[0])
            if "After drying" in desc and "W2" not in desc:
                joined = " ".join(_cell_text(c) for c in row.cells[1:])
                assert joined.strip() not in ("g", "g g")
                break
        else:
            pytest.fail("Expected unfilled After drying row in moisture worksheet")

    def test_worksheet_section_numbering_sequential(self):
        sample = _sample(tests_json=json.dumps(["appearance", "moisture"]))
        results = [
            TestResultRow(
                test_key="appearance",
                test_name="Appearance",
                method="",
                unit="",
                inputs={},
                result_value="Normal",
                result_numeric=None,
            ),
            *self._jaggery_results(),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        section_titles: list[str] = []
        seen_obs = False
        for child in doc.element.body:
            if child.tag.endswith("p"):
                text = next(
                    (p.text or "").strip()
                    for p in doc.paragraphs
                    if p._element is child
                )
                if not text:
                    continue
                if "observation table" in text.lower():
                    seen_obs = True
                    continue
                if not seen_obs or text == "Result Table:":
                    continue
                first_line = text.split("\n", 1)[0].strip()
                if first_line.endswith(":"):
                    section_titles.append(first_line)
            elif child.tag.endswith("tbl") and seen_obs:
                table = next(t for t in doc.tables if t._tbl is child)
                if (
                    len(table.rows) == 1
                    and len(table.rows[0].cells) == 1
                    and "appearance" in _cell_text(table.rows[0].cells[0]).lower()
                ):
                    section_titles.append(
                        _cell_text(table.rows[0].cells[0]).split("\n")[0]
                    )
        assert section_titles[0].startswith("1.")
        assert any(
            title.startswith("2.") and "MOISTURE" in title.upper()
            for title in section_titles
        )

    @pytest.mark.skipif(
        not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing"
    )
    def test_jaggery_table_grids_fit_nutrition_house_style(self):
        """Sauce/Jaggery protocols must not overflow after Nutrition margins."""
        sample = _sample(
            tests_json=json.dumps(["appearance", "moisture", "sulphated_ash"]),
            package_type="fssai",
            sample_name="Sauce",
            lab_code="GLG/26/555",
        )
        results = [
            TestResultRow(
                test_key="appearance",
                test_name="Appearance",
                method="",
                unit="",
                inputs={},
                result_value="Red Colored Viscous",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 73.7923, "w": 5.0677, "w2": 73.6292},
                result_value="3.22",
                result_numeric=3.22,
            ),
            TestResultRow(
                test_key="sulphated_ash",
                test_name="Sulphated ash on dry basis",
                method=TEST_CATALOG["sulphated_ash"].method,
                unit="%",
                inputs={"w1": 43.0958, "w": 5.0854, "w2": 45.114},
                result_value="41.01",
                result_numeric=41.01,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))

        ref = Document(str(NUTRITION_TEMPLATE_PATH))
        ref_summary_grid = _table_grid_col_inches(ref.tables[1])
        ref_max_width = sum(ref_summary_grid)
        canonical_width, _ = _canonical_worksheet_outer_frame()
        canonical_max_width = canonical_width / 1440
        jaggery_overflow_threshold = 7.2

        summary_idx = _summary_table_index(doc)
        assert summary_idx is not None
        summary = doc.tables[summary_idx]
        assert sum(_table_grid_col_inches(summary)) == canonical_max_width
        assert sum(_table_grid_col_inches(summary)) <= ref_max_width + 0.1
        assert sum(_table_grid_col_inches(summary)) < jaggery_overflow_threshold

        for table in doc.tables:
            grid_sum = sum(_table_grid_col_inches(table))
            if grid_sum <= 0:
                continue
            assert grid_sum == canonical_max_width, (
                f"table grid {grid_sum:.2f}in != canonical {canonical_max_width:.2f}in"
            )

        for table in doc.sections[0].header.tables:
            grid_sum = sum(_table_grid_col_inches(table))
            if grid_sum > 0:
                assert grid_sum <= ref_max_width + 0.05

    @pytest.mark.skipif(
        not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing"
    )
    def test_jaggery_table_indents_match_nutrition_reference(self):
        sample = _sample(
            tests_json=json.dumps(["appearance", "moisture"]),
            package_type="fssai",
            sample_name="Sauce",
        )
        results = [
            TestResultRow(
                test_key="appearance",
                test_name="Appearance",
                method="",
                unit="",
                inputs={},
                result_value="Red Colored Viscous",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={"w1": 73.7923, "w": 5.0677, "w2": 73.6292},
                result_value="3.22",
                result_numeric=3.22,
            ),
        ]
        out = fill_protocol_docx_bytes(sample, _header(), results)
        doc = Document(io.BytesIO(out))
        canonical_width, canonical_ind = _canonical_worksheet_outer_frame()

        summary_idx = _summary_table_index(doc)
        assert summary_idx is not None
        assert _table_tbl_ind(doc.tables[summary_idx]) == canonical_ind

        worksheet = _worksheet_table_with_text(doc, "stainless steel dish")
        assert worksheet is not None
        assert _table_tbl_ind(worksheet) == canonical_ind

        assert _table_tbl_ind(doc.tables[0]) == canonical_ind
        assert _table_grid_sum_twips(doc.tables[0]) == canonical_width

    @pytest.mark.skipif(
        not NUTRITION_TEMPLATE_PATH.exists(), reason="Nutrition template missing"
    )
    def test_jaggery_no_spacer_run_before_observation(self):
        sample = _sample(
            tests_json=json.dumps(["moisture"]),
            package_type="fssai",
        )
        out = fill_protocol_docx_bytes(sample, _header(), self._jaggery_results())
        doc = Document(io.BytesIO(out))
        assert _empty_paragraphs_between_summary_and_observation(doc) <= 1


def _row_has_cant_split(row) -> bool:
    tr_pr = row._tr.trPr
    if tr_pr is None:
        return False
    return tr_pr.find(qn("w:cantSplit")) is not None


def _paragraph_has_keep_next(paragraph: Paragraph) -> bool:
    p_pr = paragraph._element.pPr
    if p_pr is None:
        return False
    return p_pr.find(qn("w:keepNext")) is not None


def _appearance_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) == 1 and len(table.rows[0].cells) == 1:
            if "appearance" in _cell_text(table.rows[0].cells[0]).lower():
                return table
    return None


def _cell_run_font(cell) -> tuple[str | None, str | None, bool]:
    """Return (font_name, sz_half_points, bold) from the first run in a cell."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            r_pr = run._element.rPr
            if r_pr is None:
                continue
            fonts = r_pr.find(qn("w:rFonts"))
            name = fonts.get(qn("w:ascii")) if fonts is not None else None
            sz_el = r_pr.find(qn("w:sz"))
            size = sz_el.get(qn("w:val")) if sz_el is not None else None
            bold = r_pr.find(qn("w:b")) is not None
            return name, size, bold
    return None, None, False


def _paragraph_run_font(paragraph: Paragraph) -> tuple[str | None, str | None, bool]:
    for run in paragraph.runs:
        r_pr = run._element.rPr
        if r_pr is None:
            continue
        fonts = r_pr.find(qn("w:rFonts"))
        name = fonts.get(qn("w:ascii")) if fonts is not None else None
        sz_el = r_pr.find(qn("w:sz"))
        size = sz_el.get(qn("w:val")) if sz_el is not None else None
        bold = r_pr.find(qn("w:b")) is not None
        return name, size, bold
    return None, None, False


def _cell_has_no_wrap(cell) -> bool:
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return False
    return tc_pr.find(qn("w:noWrap")) is not None


def _has_blank_line_before_result_table(doc: Document) -> bool:
    if not doc.tables:
        return False
    next_el = doc.tables[0]._tbl.getnext()
    if next_el is None or not next_el.tag.endswith("p"):
        return False
    para = Paragraph(next_el, doc)
    if not (para.text or "").strip():
        following = next_el.getnext()
        if following is not None and following.tag.endswith("p"):
            return (Paragraph(following, doc).text or "").strip() == "Result Table:"
    if (para.text or "").strip() == "Result Table:":
        return False
    return False


@pytest.mark.skipif(
    not JAGGERY_TEMPLATE_PATH.exists() or not NUTRITION_TEMPLATE_PATH.exists(),
    reason="Templates missing",
)
class TestSauceProtocolLayout:
    """House-style layout: appearance box, headers, alignment, unsplit worksheets."""

    def _sauce_results(self) -> list[TestResultRow]:
        return [
            TestResultRow(
                test_key="appearance",
                test_name="Appearance",
                method="",
                unit="",
                inputs={},
                result_value="Red Colored Viscous",
                result_numeric=None,
            ),
            TestResultRow(
                test_key="moisture",
                test_name="Moisture",
                method=TEST_CATALOG["moisture"].method,
                unit="%",
                inputs={
                    "empty_dish": 68.7246,
                    "w1": 73.7923,
                    "w": 5.0677,
                    "w2": 73.6292,
                },
                result_value="3.22",
                result_numeric=3.22,
            ),
            TestResultRow(
                test_key="sulphated_ash",
                test_name="Sulphated ash on dry basis",
                method=TEST_CATALOG["sulphated_ash"].method,
                unit="%",
                inputs={
                    "w1": 43.0958,
                    "w": 5.0854,
                    "w2": 45.114,
                },
                result_value="41.01",
                result_numeric=41.01,
            ),
        ]

    def _filled_sauce_doc(self) -> Document:
        sample = _sample(
            sample_name="Sauce",
            lab_code="GLG/26/555",
            tests_json=json.dumps(["appearance", "moisture", "sulphated_ash"]),
            package_type="fssai",
        )
        out = fill_protocol_docx_bytes(
            sample,
            _header(),
            self._sauce_results(),
        )
        return Document(io.BytesIO(out))

    def test_appearance_stays_bordered_table(self):
        doc = self._filled_sauce_doc()
        app_table = _appearance_table(doc)
        assert app_table is not None
        assert "Red Colored Viscous" in _cell_text(app_table.rows[0].cells[0])
        for paragraph in doc.paragraphs:
            if "red colored viscous" in (paragraph.text or "").lower():
                pytest.fail("Appearance promoted to body paragraph")

    def test_moisture_worksheet_third_header_is_readings(self):
        doc = self._filled_sauce_doc()
        moisture = _worksheet_table_with_text(doc, "stainless steel dish")
        assert moisture is not None
        third_hdr = _cell_text(moisture.rows[0].cells[2])
        assert third_hdr == "Readings"

    def test_sulphated_ash_worksheet_third_header_is_readings(self):
        doc = self._filled_sauce_doc()
        ash = _worksheet_table_with_text(doc, "Sulphated ash")
        assert ash is not None
        third_hdr = _cell_text(ash.rows[0].cells[2])
        assert third_hdr == "Readings"

    def test_worksheet_rows_have_cant_split(self):
        doc = self._filled_sauce_doc()
        moisture = _worksheet_table_with_text(doc, "stainless steel dish")
        assert moisture is not None
        assert all(_row_has_cant_split(row) for row in moisture.rows)

    def test_worksheet_section_title_has_keep_next(self):
        doc = self._filled_sauce_doc()
        moisture = _worksheet_table_with_text(doc, "stainless steel dish")
        assert moisture is not None
        prev = moisture._tbl.getprevious()
        assert prev is not None and prev.tag.endswith("p")
        para = Paragraph(prev, doc)
        assert "MOISTURE" in (para.text or "").upper()
        assert _paragraph_has_keep_next(para)

    def test_footer_table_aligns_with_sample_table(self):
        doc = self._filled_sauce_doc()
        canonical_width, canonical_ind = _canonical_worksheet_outer_frame()
        sample = doc.tables[0]
        footer = doc.sections[0].footer.tables[0]
        assert _table_tbl_ind(sample) == canonical_ind
        assert _table_tbl_ind(footer) == canonical_ind
        assert _table_grid_sum_twips(sample) == canonical_width
        assert _table_grid_sum_twips(footer) == canonical_width

    def test_section_header_matches_nutrition_spacing_structure(self):
        doc = self._filled_sauce_doc()
        ref = Document(str(NUTRITION_TEMPLATE_PATH))
        ref_hdr = ref.sections[0].header._element
        ref_trailing = 0
        seen_tbl = False
        for child in ref_hdr:
            if child.tag.endswith("tbl"):
                seen_tbl = True
                ref_trailing = 0
            elif seen_tbl and child.tag.endswith("p"):
                ref_trailing += 1

        hdr_el = doc.sections[0].header._element
        children = list(hdr_el)
        assert children[0].tag.endswith("tbl")
        assert "Protocol No" in doc.sections[0].header.tables[0].rows[0].cells[0].text
        trailing = [c for c in children[1:] if c.tag.endswith("p")]
        assert len(trailing) == ref_trailing == 2
        for para_el in trailing:
            text = "".join(t.text or "" for t in para_el.findall(".//" + qn("w:t")))
            assert not text.strip()

    def test_all_outer_boxes_share_width_and_indent(self):
        doc = self._filled_sauce_doc()
        canonical_width, canonical_ind = _canonical_worksheet_outer_frame()
        outer_tables = [doc.tables[0], doc.tables[1], _appearance_table(doc)]
        outer_tables.extend(
            t
            for t in doc.tables
            if "description" in _cell_text(t.rows[0].cells[0]).lower()
        )
        outer_tables.append(doc.sections[0].header.tables[0])
        outer_tables.append(doc.sections[0].footer.tables[0])

        for table in outer_tables:
            assert table is not None
            assert _table_tbl_ind(table) == canonical_ind
            assert _table_grid_sum_twips(table) == canonical_width
            assert _table_outer_right_edge_twips(table) == int(canonical_ind[0]) + canonical_width

    def test_appearance_value_on_single_line(self):
        doc = self._filled_sauce_doc()
        app_table = _appearance_table(doc)
        assert app_table is not None
        text = _cell_text(app_table.rows[0].cells[0])
        assert "APPEARANCE:" in text.upper()
        assert "Red Colored Viscous" in text
        assert "\n" not in text.strip()

    def test_protocol_header_uses_cambria_10pt(self):
        doc = self._filled_sauce_doc()
        header_table = doc.sections[0].header.tables[0]
        for cell in header_table.rows[0].cells:
            name, size, bold = _cell_run_font(cell)
            if name is None and size is None:
                continue
            assert name == "Cambria"
            assert size == "20"
            assert not bold

    def test_protocol_header_cells_single_line(self):
        doc = self._filled_sauce_doc()
        header_table = doc.sections[0].header.tables[0]
        for cell in header_table.rows[0].cells:
            assert "\n" not in _cell_text(cell)
            assert _cell_has_no_wrap(cell)

    def test_blank_line_before_result_table(self):
        doc = self._filled_sauce_doc()
        assert _has_blank_line_before_result_table(doc)

    def test_result_table_title_uses_cambria_bold_13pt(self):
        doc = self._filled_sauce_doc()
        for paragraph in doc.paragraphs:
            if (paragraph.text or "").strip() == "Result Table:":
                assert paragraph.text == "Result Table:"
                name, size, bold = _paragraph_run_font(paragraph)
                assert name == "Cambria"
                assert size == "26"
                assert bold
                return
        pytest.fail("Result Table: paragraph not found")

    def test_result_table_header_row_bold_11pt(self):
        doc = self._filled_sauce_doc()
        summary = doc.tables[1]
        for cell in summary.rows[0].cells:
            name, size, bold = _cell_run_font(cell)
            assert name == "Cambria"
            assert size == "22"
            assert bold

    def test_result_table_data_cells_all_10pt(self):
        doc = self._filled_sauce_doc()
        summary = doc.tables[1]
        for row in summary.rows[1:]:
            for cell in row.cells:
                name, size, bold = _cell_run_font(cell)
                if name is None and size is None:
                    continue
                assert name == "Cambria"
                assert size == "20"
                assert not bold

    def test_worksheet_data_cells_all_10pt(self):
        doc = self._filled_sauce_doc()
        moisture = _worksheet_table_with_text(doc, "stainless steel dish")
        assert moisture is not None
        desc_cell = moisture.rows[1].cells[0]
        reading_cell = moisture.rows[1].cells[1]
        for cell in (desc_cell, reading_cell):
            name, size, bold = _cell_run_font(cell)
            assert name == "Cambria"
            assert size == "20"
            assert not bold

    def test_sauce_layout_preview_written(self):
        doc_bytes = fill_protocol_docx_bytes(
            _sample(
                sample_name="Sauce",
                lab_code="GLG/26/555",
                tests_json=json.dumps(["appearance", "moisture", "sulphated_ash"]),
                package_type="fssai",
            ),
            _header(),
            self._sauce_results(),
        )
        preview_docx = PROJECT_ROOT / "tmp" / "Protocol_Sauce_layout_preview.docx"
        preview_docx.parent.mkdir(parents=True, exist_ok=True)
        preview_docx.write_bytes(doc_bytes)
        assert preview_docx.stat().st_size > 1000
