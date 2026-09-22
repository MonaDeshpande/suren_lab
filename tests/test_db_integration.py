"""
DB integration tests: Reception persistence, Analyst protocol fields, CTR symbols.

Run: pytest -m integration tests/test_db_integration.py
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from docx import Document
from pypdf import PdfReader

from db.connection import get_db
from services.appearance_master import get_or_create_appearance, list_appearances
from services.customers import Customer
from services.docx_filler import fill_docx_bytes, suggest_docx_filename
from services.pdf_generator import generate_pdf_bytes, suggest_pdf_filename
from services.protocol_store import (
    ProtocolHeader,
    format_analysis_date_range,
    get_protocol_header,
    upsert_protocol_header,
    validate_analysis_date_range,
)
from services.protocols.test_catalog import (
    CATEGORY_CATTLE_FEED_FERTILIZER,
    CATEGORY_FOOD,
    CATEGORY_MICRO,
    CATEGORY_WATER,
    MICRO_TEST_KEYS,
    WATER_MICRO_TEST_KEYS,
    WATER_TEST_KEYS,
    normalize_category,
)
from services.requests import (
    SampleRow,
    TestRequestData,
    get_test_request,
    save_test_request,
    storage_temperature_for_save,
    sync_verify_sample_code,
)
from services.samples import get_by_code
from tests.conftest import cleanup_test_data, sample_verification_kwargs

pytestmark = pytest.mark.integration

PREVIEW_LAB_CODE = "SLS/26/900"


def _preview_customer(gst: str) -> Customer:
    return Customer(
        customer_name="ABC Foods Pvt Ltd",
        address="Plot 12, MIDC, Nashik, Maharashtra",
        contact_person="Mr. Ramesh Patil",
        contact_number="9876543210",
        email="ramesh@abcfoods.in",
        gst_number=gst,
    )


def _verification(lab_code: str) -> dict:
    return sample_verification_kwargs(verify_lab_code=lab_code)


def build_multi_type_ctr_data(
    chem_analyst_id: int,
    micro_analyst_id: int,
    *,
    gst: str,
    lab_code: str = PREVIEW_LAB_CODE,
) -> TestRequestData:
    """TC-REC-037 payload: one customer, mixed categories + logo formats."""
    jaggery_logo = ["moisture", "total_ash", "acid_insoluble_ash", "added_color"]
    jaggery_nologo = [
        "invert_sugar",
        "reducing_sugar",
        "sucrose",
        "sulphated_ash",
        "sulphur_dioxide",
    ]
    verify = _verification(lab_code)
    return TestRequestData(
        customer=_preview_customer(gst),
        request_date=date(2026, 7, 27),
        lab_code=lab_code,
        number_of_samples=5,
        sampling_by_lab=True,
        storage_temperature=storage_temperature_for_save("4°C", ""),
        decision_rule=True,
        service_type="regular",
        delivery_mode="Collect, Email/Whatsapp",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Potable Water",
                batch_code="---",
                quantity="1 lit",
                parameters="Drinking Water",
                category=CATEGORY_WATER,
                test_keys=list(WATER_TEST_KEYS) + list(WATER_MICRO_TEST_KEYS),
                assigned_analyst_id=chem_analyst_id,
                assigned_micro_analyst_id=micro_analyst_id,
                protocol_no="P-W-DB-001",
                storage_temperature=storage_temperature_for_save("4°C", ""),
                **verify,
            ),
            SampleRow(
                sr_no=2,
                sample_name="Cooked Food: Paneer Gravy",
                quantity="250 gm",
                parameters="Paneer Gravy",
                category=CATEGORY_MICRO,
                test_keys=list(MICRO_TEST_KEYS),
                assigned_analyst_id=chem_analyst_id,
                protocol_no="P-M-DB-001",
                storage_temperature=storage_temperature_for_save("2°C to 8°C", ""),
                **verify,
            ),
            SampleRow(
                sr_no=3,
                sample_name="Jaggery",
                batch_code="B-01",
                quantity="500 g",
                parameters="FSSAI",
                parameters_select="FSSAI",
                package_type="fssai",
                category=CATEGORY_FOOD,
                test_keys=["moisture"],
                tests_with_logo=["moisture"],
                report_format="with_logo",
                assigned_analyst_id=chem_analyst_id,
                protocol_no="P-F-DB-001",
                storage_temperature=storage_temperature_for_save("Room Temp", ""),
                **verify,
            ),
            SampleRow(
                sr_no=4,
                sample_name="Jaggery",
                batch_code="B-02",
                quantity="900 g",
                parameters="FSSAI",
                parameters_select="FSSAI",
                package_type="fssai",
                category=CATEGORY_FOOD,
                test_keys=jaggery_logo + jaggery_nologo,
                tests_with_logo=jaggery_logo,
                tests_without_logo=jaggery_nologo,
                report_format="both",
                assigned_analyst_id=chem_analyst_id,
                protocol_no="P-J-DB-001",
                storage_temperature=storage_temperature_for_save("Other", "25"),
                **verify,
            ),
            SampleRow(
                sr_no=5,
                sample_name="Cattle Feed Mix",
                quantity="1 kg",
                parameters="Standard cattle feed panel",
                category=CATEGORY_CATTLE_FEED_FERTILIZER,
                assigned_analyst_id=chem_analyst_id,
                protocol_no="P-CF-DB-001",
                storage_temperature=storage_temperature_for_save("Other", "18"),
                **verify,
            ),
        ],
    )


@pytest.fixture
def test_request_bundle(require_db, qa_analyst_pair, db_test_gst):
    """Persist multi-type CTR; teardown removes all rows for the test GST."""
    chem, micro = qa_analyst_pair
    data = build_multi_type_ctr_data(chem.id, micro.id, gst=db_test_gst)
    saved = save_test_request(data, actor=None)
    assert saved.request_id is not None
    yield saved
    cleanup_test_data(gst=db_test_gst)


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


class TestCategoryPerSample:
    def test_category_persisted_per_sample_row(self, test_request_bundle):
        loaded = get_test_request(test_request_bundle.request_id)
        assert loaded is not None
        by_sr = {s.sr_no: s for s in loaded.samples}
        assert normalize_category(by_sr[1].category) == CATEGORY_WATER
        assert normalize_category(by_sr[2].category) == CATEGORY_MICRO
        assert normalize_category(by_sr[3].category) == CATEGORY_FOOD
        assert normalize_category(by_sr[4].category) == CATEGORY_FOOD
        assert (
            normalize_category(by_sr[5].category) == CATEGORY_CATTLE_FEED_FERTILIZER
        )

        water = get_by_code(by_sr[1].sample_code)
        assert water is not None
        assert water.category == CATEGORY_WATER


class TestMultiSamplePayload:
    def test_multi_sample_distinct_request_details(self, test_request_bundle):
        loaded = get_test_request(test_request_bundle.request_id)
        assert loaded is not None
        by_sr = {s.sr_no: s for s in loaded.samples}

        assert by_sr[1].quantity == "1 lit"
        assert by_sr[2].quantity == "250 gm"
        assert by_sr[3].quantity == "500 g"
        assert by_sr[1].protocol_no == "P-W-DB-001"
        assert by_sr[2].protocol_no == "P-M-DB-001"
        assert by_sr[3].protocol_no == "P-F-DB-001"
        assert by_sr[4].report_format == "both"
        assert by_sr[3].report_format == "with_logo"
        assert "moisture" in by_sr[3].tests_with_logo
        assert "invert_sugar" in by_sr[4].tests_without_logo


class TestStorageTemperature:
    def test_storage_temperature_persisted_with_celsius(self, test_request_bundle):
        loaded = get_test_request(test_request_bundle.request_id)
        assert loaded is not None
        by_sr = {s.sr_no: s for s in loaded.samples}

        assert by_sr[1].storage_temperature == "4°C"
        assert by_sr[2].storage_temperature == "2°C to 8°C"
        assert by_sr[3].storage_temperature == "Room Temp"
        assert by_sr[4].storage_temperature == "25°C"
        assert by_sr[5].storage_temperature == "18°C"
        assert loaded.storage_temperature == "4°C"


class TestVerifySampleCode:
    def test_verify_sample_code_auto_filled_on_save(self, test_request_bundle):
        session: dict = {}
        for sample in test_request_bundle.samples:
            assert sample.sample_code
            synced = sync_verify_sample_code(
                sample.sample_code,
                session=session,
                sr_no=sample.sr_no,
            )
            assert synced == sample.sample_code
            assert session[f"verify_sample_code_{sample.sr_no}"] == sample.sample_code

            record = get_by_code(sample.sample_code)
            assert record is not None
            assert record.sample_code == sample.sample_code

        with get_db() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT sample_code
                      FROM request_samples
                     WHERE request_id = %s
                     ORDER BY sr_no
                    """,
                    (test_request_bundle.request_id,),
                )
                codes = [row[0] for row in cur.fetchall()]
        assert len(codes) == 5
        assert codes[0] == "SLS/26/900/01"


class TestCheckmarkSymbols:
    def test_ctr_checklist_uses_checkmark_not_x(self, test_request_bundle):
        loaded = get_test_request(test_request_bundle.request_id)
        assert loaded is not None

        pdf = generate_pdf_bytes(loaded)
        text = _pdf_text(pdf)
        assert "Sample Verification Checklist" in text
        assert "✓" in text
        assert "Yes ( X )" not in text

        docx = fill_docx_bytes(loaded)
        doc = Document(io.BytesIO(docx))
        body = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    body += cell.text
        assert "✓" in body
        assert "Yes ( X )" not in body


class TestAppearanceMaster:
    def test_appearance_master_create_and_list(
        self,
        require_db,
        test_request_bundle,
        appearance_test_cleanup,
    ):
        label = f"DB integration creamy {uuid.uuid4().hex[:6]}"
        created = get_or_create_appearance(label)
        appearance_test_cleanup.append(created.id)

        texts = {o.appearance_text for o in list_appearances()}
        assert label in texts

        sample = test_request_bundle.samples[0]
        assert sample.id is not None
        upsert_protocol_header(
            ProtocolHeader(
                sample_id=sample.id,
                protocol_no=sample.protocol_no,
                issued_to="QA Analyst",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 27),
                appearance_text=label,
            ),
            actor=None,
        )
        header = get_protocol_header(sample.id)
        assert header is not None
        assert header.appearance_text == label


class TestAnalysisDateRange:
    def test_analysis_date_range_persisted_and_formatted(
        self,
        require_db,
        test_request_bundle,
    ):
        sample = test_request_bundle.samples[0]
        assert sample.id is not None

        date_from = date(2026, 7, 1)
        date_to = date(2026, 7, 10)
        assert validate_analysis_date_range(date_from, date_to) == []
        assert validate_analysis_date_range(date_to, date_from) != []

        upsert_protocol_header(
            ProtocolHeader(
                sample_id=sample.id,
                protocol_no=sample.protocol_no,
                issued_to="QA Analyst",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 27),
                date_of_analysis_from=date_from,
                date_of_analysis_to=date_to,
            ),
            actor=None,
        )
        header = get_protocol_header(sample.id)
        assert header is not None
        assert header.date_of_analysis_from == date_from
        assert header.date_of_analysis_to == date_to
        assert format_analysis_date_range(date_from, date_to) == (
            "01/07/2026 to 10/07/2026"
        )
