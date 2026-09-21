"""
Analyst workflow + protocol preview tests.

Run: pytest tests/test_analyst_workflow.py -v -s
Integration tests require PostgreSQL.
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from services.customers import Customer
from services.e2e_app_flow import (
    E2EScenario,
    _ctr_request,
    build_demo_scenarios,
    prepare_app_e2e,
    run_analyst,
)
from services.protocol_docx import fill_protocol_docx_bytes
from services.protocol_store import (
    ProtocolHeader,
    default_protocol_disclaimer_text,
    get_protocol_header,
    list_results,
    save_test_result,
    upsert_protocol_header,
)
from services.protocol_pdf import generate_protocol_documents
from services.protocols.test_catalog import (
    MICRO_TEST_KEYS,
    WATER_MICRO_TEST_KEYS,
    default_water_micro_procedure,
)
from services.requests import SampleRow, TestRequestData, save_test_request
from services.samples import (
    REPORT_FORMAT_BOTH,
    REPORT_FORMAT_WITHOUT_LOGO,
    REPORT_FORMAT_WITH_LOGO,
    SampleRecord,
    assigned_test_keys_for_sample,
    get_by_code,
    logo_test_keys,
    no_logo_test_keys,
)
from tests.conftest import cleanup_test_data, sample_verification_kwargs

OUTPUT_ROOT = Path(__file__).resolve().parent / "output_preview" / "protocol"


@pytest.fixture(scope="session")
def protocol_output_root() -> Path:
    for sub in ("water", "micro", "food"):
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    return OUTPUT_ROOT


def _docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _header_has_logo(doc: Document) -> bool:
    for section in doc.sections:
        if section.header._element.findall(".//" + qn("w:drawing")):
            return True
    return False


def _write_protocol_preview(
    out_dir: Path,
    *,
    filename: str,
    docx_bytes: bytes,
    pdf_bytes: bytes | None,
) -> tuple[Path, Path | None]:
    safe = filename.replace("/", "-").replace("\\", "-")
    docx_path = out_dir / safe
    docx_path.write_bytes(docx_bytes)
    pdf_path = None
    if pdf_bytes:
        pdf_path = out_dir / (Path(safe).stem + ".pdf")
        pdf_path.write_bytes(pdf_bytes)
    return docx_path, pdf_path


def _food_sample_record(**overrides) -> SampleRecord:
    base = dict(
        id=1,
        request_id=1,
        sample_code="SLS/26/562/01",
        sr_no=1,
        sample_name="Jaggery",
        batch_code="B-01",
        quantity="500 g",
        parameters="FSSAI",
        tests_to_perform="",
        status="in_progress",
        analyst_remarks="",
        created_at=None,
        expires_at=None,
        category="food",
        tests_json=json.dumps(["moisture"]),
        lab_code="SLS/26/562",
        report_format=REPORT_FORMAT_WITH_LOGO,
        tests_with_logo_json=json.dumps(["moisture"]),
        tests_without_logo_json="[]",
        package_type="fssai",
    )
    base.update(overrides)
    return SampleRecord(**base)


def _jaggery_header(**overrides) -> ProtocolHeader:
    base = dict(
        sample_id=1,
        protocol_no="P-J-001",
        issued_to="Analyst",
        issued_by="Reception",
        sample_received_on=date(2026, 7, 7),
        date_of_analysis=date(2026, 7, 14),
        appearance_text="Brown colored solid jaggery.",
        protocol_disclaimer_text=default_protocol_disclaimer_text(),
    )
    base.update(overrides)
    return ProtocolHeader(**base)


def _reception_save(scenario: E2EScenario) -> str:
    """Persist reception data without validate_request (e2e demo GST lengths)."""
    data = _ctr_request(
        customer=scenario.customer,
        lab_code=scenario.lab_code,
        sample=scenario.sample,
        request_date=scenario.request_date,
    )
    saved = save_test_request(data, actor=None)
    code = (saved.samples[0].sample_code or "").strip()
    if not code:
        raise RuntimeError("Reception save did not assign a sample code")
    return code


def test_analyst_screen_branding_logo():
    """With-logo vs without-logo food samples expose different assigned key sets."""
    both = _food_sample_record(
        report_format=REPORT_FORMAT_BOTH,
        tests_json=json.dumps(["moisture", "total_ash"]),
        tests_with_logo_json=json.dumps(["moisture"]),
        tests_without_logo_json=json.dumps(["total_ash"]),
    )
    assert logo_test_keys(both) == {"moisture"}
    assert no_logo_test_keys(both) == {"total_ash"}
    assigned = set(assigned_test_keys_for_sample(both))
    assert assigned == {"moisture", "total_ash"}

    with_logo = _food_sample_record(report_format=REPORT_FORMAT_WITH_LOGO)
    without_logo = _food_sample_record(report_format=REPORT_FORMAT_WITHOUT_LOGO)
    doc_logo = Document(
        io.BytesIO(
            fill_protocol_docx_bytes(
                with_logo,
                _jaggery_header(),
                [],
            )
        )
    )
    doc_plain = Document(
        io.BytesIO(
            fill_protocol_docx_bytes(
                without_logo,
                _jaggery_header(),
                [],
            )
        )
    )
    assert _header_has_logo(doc_logo)
    assert not _header_has_logo(doc_plain)


@pytest.mark.integration
def test_formula_evaluation_and_report_reflection(require_db, qa_analyst_pair, db_test_gst):
    """Moisture worksheet calculation persists and appears in protocol output."""
    chem, _micro = qa_analyst_pair
    data = TestRequestData(
        customer=Customer(
            customer_name="Formula QA Co",
            address="Test",
            contact_person="QA",
            contact_number="9999999999",
            email="qa@example.com",
            gst_number=db_test_gst,
        ),
        request_date=date(2026, 7, 17),
        lab_code="SLS/26/301",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Jaggery",
                parameters="FSSAI",
                package_type="fssai",
                test_keys=["moisture"],
                assigned_analyst_id=chem.id,
                protocol_no="P-FORM-001",
                **sample_verification_kwargs(verify_lab_code="SLS/26/301"),
            )
        ],
    )
    saved = save_test_request(data, actor=None)
    sample = get_by_code(saved.samples[0].sample_code)
    assert sample is not None
    try:
        upsert_protocol_header(
            ProtocolHeader(
                sample_id=sample.id,
                protocol_no="P-FORM-001",
                issued_to="QA Analyst",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 17),
                date_of_analysis=date(2026, 7, 18),
                appearance_text="Brown",
            ),
            actor=None,
        )
        row = save_test_result(
            sample.id,
            "moisture",
            {"empty_dish": "50.0", "w1": "55.0", "w": "5.0", "w2": "54.8235"},
            actor=None,
        )
        assert row.result_numeric is not None
        assert float(row.result_numeric) == pytest.approx(3.53, rel=1e-2)

        header = get_protocol_header(sample.id)
        assert header is not None
        results = list_results(sample.id)
        docx = fill_protocol_docx_bytes(sample, header, results)
        text = _docx_text(docx)
        assert "3.53" in text or "3.5" in text
    finally:
        cleanup_test_data(gst=db_test_gst)


@pytest.mark.integration
def test_multi_sample_partial_execution(require_db, qa_analyst_pair, db_test_gst):
    """Two samples under one customer keep independent saved results."""
    chem, micro = qa_analyst_pair
    verify = sample_verification_kwargs(verify_lab_code="SLS/26/800")
    data = TestRequestData(
        customer=Customer(
            customer_name="Multi Sample QA",
            address="",
            contact_person="QA",
            contact_number="8888888888",
            email="multi@example.com",
            gst_number=db_test_gst,
        ),
        request_date=date(2026, 7, 20),
        lab_code="SLS/26/800",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Sample A",
                category="food",
                parameters="FSSAI",
                package_type="fssai",
                test_keys=["moisture"],
                assigned_analyst_id=chem.id,
                protocol_no="P-A-001",
                **verify,
            ),
            SampleRow(
                sr_no=2,
                sample_name="Sample B",
                category="micro",
                test_keys=list(MICRO_TEST_KEYS),
                assigned_analyst_id=chem.id,
                protocol_no="P-B-001",
                **verify,
            ),
        ],
    )
    saved = save_test_request(data, actor=None)
    code_a = saved.samples[0].sample_code
    code_b = saved.samples[1].sample_code
    sample_a = get_by_code(code_a)
    sample_b = get_by_code(code_b)
    assert sample_a is not None and sample_b is not None
    try:
        upsert_protocol_header(
            ProtocolHeader(sample_id=sample_a.id, protocol_no="P-A-001"),
            actor=None,
        )
        upsert_protocol_header(
            ProtocolHeader(sample_id=sample_b.id, protocol_no="P-B-001"),
            actor=None,
        )
        save_test_result(
            sample_a.id,
            "moisture",
            {"empty_dish": "50", "w1": "55", "w": "5", "w2": "54.8"},
            actor=None,
        )
        save_test_result(
            sample_b.id,
            "e_coli",
            {
                "result_value": "Absent",
                "result_unit": "cfu/gm",
                "method_override": "Custom micro method B",
            },
            actor=None,
        )

        results_a = list_results(sample_a.id)
        results_b = list_results(sample_b.id)
        assert {r.test_key for r in results_a} == {"moisture"}
        assert {r.test_key for r in results_b} == {"e_coli"}
        assert results_b[0].method == "Custom micro method B"
        assert all(r.test_key != "e_coli" for r in results_a)
    finally:
        cleanup_test_data(gst=db_test_gst)


@pytest.mark.integration
def test_micro_sample_custom_unit_and_method_editing(
    require_db, qa_analyst_pair, db_test_gst
):
    chem, _micro = qa_analyst_pair
    data = TestRequestData(
        customer=Customer(
            customer_name="Micro QA Co",
            address="",
            contact_person="QA",
            contact_number="7777777777",
            email="micro@example.com",
            gst_number=db_test_gst,
        ),
        request_date=date(2026, 7, 4),
        lab_code="SLS/26/546",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Cooked Food: Paneer Gravy",
                quantity="250 gm",
                parameters="Paneer Gravy",
                category="micro",
                test_keys=list(MICRO_TEST_KEYS),
                assigned_analyst_id=chem.id,
                protocol_no="P-M-001",
                **sample_verification_kwargs(verify_lab_code="SLS/26/546"),
            )
        ],
    )
    saved = save_test_request(data, actor=None)
    sample = get_by_code(saved.samples[0].sample_code)
    assert sample is not None
    try:
        upsert_protocol_header(
            ProtocolHeader(
                sample_id=sample.id,
                protocol_no="P-M-001",
                issued_to="QA Analyst",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 4),
                date_of_analysis=date(2026, 7, 9),
            ),
            actor=None,
        )
        save_test_result(
            sample.id,
            "total_plate_count",
            {
                "result_value": "2.0 x 10²",
                "result_unit": "cfu/100ml",
                "method_override": "Custom IS 5402:2018 Rev A",
            },
            actor=None,
        )
        save_test_result(
            sample.id,
            "e_coli",
            {
                "result_value": "Inconclusive",
                "result_unit": "cfu/gm",
                "method_override": "Lab SOP-M-2026",
            },
            actor=None,
        )
        results = list_results(sample.id)
        by_key = {r.test_key: r for r in results}
        assert by_key["total_plate_count"].unit == "cfu/100ml"
        assert by_key["total_plate_count"].method == "Custom IS 5402:2018 Rev A"
        assert by_key["e_coli"].result_value == "Inconclusive"

        header = get_protocol_header(sample.id)
        assert header is not None
        docx_bytes, pdf_bytes, docx_name, _, _ = generate_protocol_documents(
            sample, header, results
        )
        text = _docx_text(docx_bytes)
        assert "cfu/100ml" in text
        assert "Custom IS 5402:2018 Rev A" in text
        assert "Inconclusive" in text
        assert len(docx_bytes) > 5000
    finally:
        cleanup_test_data(gst=db_test_gst)


@pytest.mark.integration
def test_editable_disclaimer_persistence(require_db, qa_analyst_pair, db_test_gst):
    chem, _micro = qa_analyst_pair
    data = TestRequestData(
        customer=Customer(
            customer_name="Disclaimer QA",
            address="",
            contact_person="QA",
            contact_number="6666666666",
            email="disc@example.com",
            gst_number=db_test_gst,
        ),
        request_date=date(2026, 7, 7),
        lab_code="SLS/26/563",
        samples=[
            SampleRow(
                sr_no=1,
                sample_name="Jaggery",
                parameters="FSSAI",
                package_type="fssai",
                test_keys=["moisture"],
                assigned_analyst_id=chem.id,
                protocol_no="P-DISC-001",
                **sample_verification_kwargs(verify_lab_code="SLS/26/563"),
            )
        ],
    )
    saved = save_test_request(data, actor=None)
    sample = get_by_code(saved.samples[0].sample_code)
    assert sample is not None
    edited_line = (
        "4. Sample stored for two Weeks and test Report for eighteen months "
        "from the date received."
    )
    disclaimer = default_protocol_disclaimer_text().replace(
        "4. Sample stored for one Week and test Report for one year from the date received.",
        edited_line,
    )
    try:
        upsert_protocol_header(
            ProtocolHeader(
                sample_id=sample.id,
                protocol_no="P-DISC-001",
                issued_to="QA Analyst",
                issued_by="Reception",
                sample_received_on=date(2026, 7, 7),
                date_of_analysis=date(2026, 7, 14),
                appearance_text="Brown",
                protocol_disclaimer_text=disclaimer,
            ),
            actor=None,
        )
        save_test_result(
            sample.id,
            "moisture",
            {"empty_dish": "50", "w1": "55", "w": "5", "w2": "54.8235"},
            actor=None,
        )
        header = get_protocol_header(sample.id)
        results = list_results(sample.id)
        docx_bytes, pdf_bytes, _, _, _ = generate_protocol_documents(
            sample, header, results
        )
        assert "two Weeks" in _docx_text(docx_bytes)
        if pdf_bytes:
            assert "two Weeks" in _pdf_text(pdf_bytes)
    finally:
        cleanup_test_data(gst=db_test_gst)


@pytest.mark.integration
def test_water_protocol_parameters_and_micro_final_page(
    require_db, qa_analyst_pair, db_test_gst, protocol_output_root
):
    chem, micro = qa_analyst_pair
    prepare_app_e2e()
    scenarios = build_demo_scenarios(chem.id, micro.id)
    water = next(s for s in scenarios if s.label == "01_water_abc_foods")
    cleanup_test_data(gst=water.customer.gst_number)
    sample_code = _reception_save(water)
    custom_ecoli_proc = "Custom E.coli membrane filter procedure for QA validation."
    inputs = dict(water.analyst_inputs_by_key)
    inputs["water_e_coli"] = {
        "procedure": custom_ecoli_proc,
        "result_obs": "absent",
    }
    water = replace(water, analyst_inputs_by_key=inputs)
    run_analyst(sample_code, water, chem_analyst_id=chem.id, micro_analyst_id=micro.id)
    sample = get_by_code(sample_code)
    assert sample is not None
    header = get_protocol_header(sample.id)
    results = list_results(sample.id)
    docx_bytes, pdf_bytes, docx_name, pdf_name, _ = generate_protocol_documents(
        sample, header, results
    )
    text = _docx_text(docx_bytes)
    assert "OBSERVATION TABLE" in text.upper()
    assert custom_ecoli_proc in text
    assert "7.320" in text or "7.32" in text

    out_dir = protocol_output_root / "water"
    docx_path, pdf_path = _write_protocol_preview(
        out_dir,
        filename=docx_name,
        docx_bytes=docx_bytes,
        pdf_bytes=pdf_bytes,
    )
    print(f"\nWater protocol preview: {docx_path}")
    if pdf_path:
        print(f"Water protocol PDF: {pdf_path}")
    cleanup_test_data(gst=water.customer.gst_number)


@pytest.mark.integration
def test_protocol_preview_generates_water_micro_food(
    require_db, qa_analyst_pair, protocol_output_root
):
    """Generate water, micro, and food protocol artifacts for layout review."""
    chem, micro = qa_analyst_pair
    prepare_app_e2e()
    scenarios = build_demo_scenarios(chem.id, micro.id)
    targets = {
        "01_water_abc_foods": "water",
        "02_micro_xyz_caterers": "micro",
        "03_food_nutrihealth": "food",
    }
    written: list[str] = []
    for scenario in scenarios:
        subdir = targets.get(scenario.label)
        if not subdir:
            continue
        cleanup_test_data(gst=scenario.customer.gst_number)
        sample_code = _reception_save(scenario)
        run_analyst(
            sample_code,
            scenario,
            chem_analyst_id=chem.id,
            micro_analyst_id=micro.id,
        )
        sample = get_by_code(sample_code)
        assert sample is not None
        header = get_protocol_header(sample.id)
        results = list_results(sample.id)
        assert header is not None and results
        docx_bytes, pdf_bytes, docx_name, _, _ = generate_protocol_documents(
            sample, header, results
        )
        assert len(docx_bytes) > 5000
        out_dir = protocol_output_root / subdir
        docx_path, pdf_path = _write_protocol_preview(
            out_dir,
            filename=docx_name,
            docx_bytes=docx_bytes,
            pdf_bytes=pdf_bytes,
        )
        written.append(str(docx_path))
        if pdf_path:
            written.append(str(pdf_path))
        cleanup_test_data(gst=scenario.customer.gst_number)

    assert len(written) >= 3
    print("\nProtocol previews written:")
    for path in written:
        print(f"  {path}")
